import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from pathlib import Path
files = [Path(r'C:\Users\Usuario\Downloads\Informe_S.docx'), Path(r'.\outputs\StockFacil\docs\Informe_StockFacil_PC_Avance_con_documentacion_funcional.docx')]
for f in files:
    print('\nFILE:', f, 'exists:', f.exists(), 'bytes:', f.stat().st_size if f.exists() else 0)
    if not f.exists():
        continue
    d = Document(f)
    print('paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'figures', len(d.inline_shapes), 'sections', len(d.sections))
    print('HEADINGS:')
    for i, p in enumerate(d.paragraphs):
        if p.style.name.startswith('Heading') or p.style.name == 'Title':
            print(i, p.style.name, p.text[:180])
    print('TABLE HEADERS:')
    for i, t in enumerate(d.tables):
        print(i, ' | '.join(c.text.replace('\n', ' / ')[:70] for c in t.rows[0].cells))
    print('FIRST PARAGRAPHS:')
    for i, p in enumerate(d.paragraphs[:18]):
        if p.text.strip():
            print(i, p.text[:260])
