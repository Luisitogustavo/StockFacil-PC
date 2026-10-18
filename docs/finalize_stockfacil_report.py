from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont


PATH = Path(r"C:\Users\Usuario\Documents\Codex\2026-10-04\hola-me-gustar-a-saber-si\outputs\StockFacil\docs\Informe_StockFacil_PC_Avance_con_documentacion_funcional.docx")
doc = Document(PATH)


def find_para(starts):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise ValueError(f"No se encontró el párrafo: {starts}")


def insert_para_after(anchor, text="", style="Normal", bold_prefix=None, page_break=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    elif text:
        p.add_run(text)
    if page_break:
        p.paragraph_format.page_break_before = True
    anchor_element = anchor._p if hasattr(anchor, "_p") else anchor._tbl
    anchor_element.addnext(p._p)
    return p


def insert_table_after(anchor, headers, rows, widths=None, size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "23466F")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(size)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_together = True
                for run in p.runs:
                    run.font.size = Pt(size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    anchor_element = anchor._p if hasattr(anchor, "_p") else anchor._tbl
    anchor_element.addnext(table._tbl)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_mark = OxmlElement("w:tblHeader")
    header_mark.set(qn("w:val"), "true")
    header_props.append(header_mark)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def highlight_para(p):
    for run in p.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.paragraph_format.keep_together = True


def add_hyperlink(paragraph, label, url):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_business_flow(path):
    width, height = 1600, 1960
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    body = ImageFont.truetype(regular_path, 27)
    small = ImageFont.truetype(regular_path, 23)
    bold = ImageFont.truetype(bold_path, 27)
    title = ImageFont.truetype(bold_path, 35)
    navy, blue, light, green, red, gray = "#193653", "#2D6496", "#EEF4FA", "#E6F3E9", "#FFF0EE", "#46586B"
    draw.text((width // 2, 18), "CONTROL DE ENTRADAS Y SALIDAS DE INVENTARIO", font=title, fill=navy, anchor="mt")

    def centered(box, text, font=body, fill=navy, outline=blue, color="white", radius=22):
        x1, y1, x2, y2 = box
        draw.rounded_rectangle(box, radius=radius, fill=color, outline=outline, width=4)
        lines = text.split("\n")
        heights = [draw.textbbox((0, 0), line, font=font)[3] - draw.textbbox((0, 0), line, font=font)[1] for line in lines]
        total = sum(heights) + 8 * (len(lines) - 1)
        y = y1 + (y2 - y1 - total) / 2
        for line, h in zip(lines, heights):
            draw.text(((x1 + x2) / 2, y), line, font=font, fill=fill, anchor="mt")
            y += h + 8

    def diamond(cx, cy, w, h, text):
        points = [(cx, cy-h//2), (cx+w//2, cy), (cx, cy+h//2), (cx-w//2, cy)]
        draw.polygon(points, fill="#FFF8E5", outline="#A47B18")
        draw.line(points + [points[0]], fill="#A47B18", width=4, joint="curve")
        lines = text.split("\n")
        heights = [draw.textbbox((0, 0), line, font=small)[3] - draw.textbbox((0, 0), line, font=small)[1] for line in lines]
        y = cy - (sum(heights) + 7*(len(lines)-1))/2
        for line, h in zip(lines, heights):
            draw.text((cx, y), line, font=small, fill=navy, anchor="mt")
            y += h + 7

    def arrow(points, label=None, label_at=None, color=blue):
        draw.line(points, fill=color, width=5, joint="curve")
        x1, y1 = points[-2]
        x2, y2 = points[-1]
        import math
        angle = math.atan2(y2-y1, x2-x1)
        size = 18
        left = (x2-size*math.cos(angle-math.pi/6), y2-size*math.sin(angle-math.pi/6))
        right = (x2-size*math.cos(angle+math.pi/6), y2-size*math.sin(angle+math.pi/6))
        draw.polygon([(x2, y2), left, right], fill=color)
        if label and label_at:
            draw.text(label_at, label, font=bold, fill=color, anchor="mm")

    # Main process path and the two business decisions.
    draw.ellipse((590, 75, 1010, 165), fill=green, outline=blue, width=4)
    draw.text((800, 120), "Inicio", font=bold, fill=navy, anchor="mm")
    centered((280, 215, 1320, 405), "Encargado elige el producto y registra\nsi ingresan o salen unidades, la cantidad\ny una nota de referencia opcional.")
    diamond(800, 560, 600, 250, "¿Producto activo\ny cantidad mayor que cero?")
    centered((25, 475, 310, 635), "No: corregir\no cancelar", small, "#9A342D", "#B04B42", "#FFF0EE", 18)
    diamond(800, 875, 500, 210, "¿Es una entrada\nde stock?")
    centered((55, 1120, 690, 1300), "Sí: sumar las unidades\na la existencia del producto.", body, navy, blue, light)
    centered((910, 1120, 1545, 1300), "Salida: comprobar que la cantidad\nno supere el saldo disponible.", body, navy, blue, light)
    diamond(1228, 1460, 430, 190, "¿Hay saldo\nsuficiente?")
    centered((1215, 1665, 1560, 1815), "No: informar el saldo\ny permitir corregir\no cancelar.", small, "#9A342D", "#B04B42", "#FFF0EE", 18)
    centered((300, 1500, 1030, 1680), "Si la entrada es válida o el saldo alcanza:\nguardar el movimiento y actualizar el stock.", body, navy, blue, green)
    draw.ellipse((390, 1805, 1210, 1900), fill=green, outline=blue, width=4)
    draw.text((800, 1852), "Consultar el saldo actualizado en la app y el portal", font=small, fill=navy, anchor="mm")

    arrow([(800,165),(800,215)])
    arrow([(800,405),(800,435)])
    arrow([(500,560),(310,560)], "No", (405,535), color="#B04B42")
    arrow([(800,685),(800,770)], "Sí", (845,730))
    # Correction branch returns to the entry step.
    arrow([(165,475),(165,310),(280,310)], color="#B04B42")
    arrow([(550,875),(370,875),(370,1120)], "Sí", (430,850))
    arrow([(1050,875),(1225,875),(1225,1120)], "No", (1135,850))
    arrow([(370,1300),(370,1415),(565,1415),(565,1500)])
    arrow([(1225,1300),(1225,1365),(1228,1365)])
    arrow([(1443,1460),(1510,1460),(1510,1665)], "No", (1475,1430), color="#B04B42")
    arrow([(1228,1555),(1228,1590),(1030,1590)], "Sí", (1190,1570))
    arrow([(665,1680),(665,1740),(800,1740),(800,1805)])
    draw.text((width//2, 1930), "El flujo requiere conexión con el servicio; la operación sin conexión aún no está implementada.", font=small, fill=gray, anchor="mm")
    image.save(path, "PNG")


# A visible reading key lets the student distinguish the completed blueprint from work
# that needs implementation, proof, a real business validation, or hosting.
intro_heading = find_para("Introducción")
note = doc.add_paragraph(style="Normal")
note.add_run("Nota de lectura:").bold = True
note.add_run(" este archivo es un borrador de trabajo. Las marcas amarillas señalan tareas que requieren completar código, incorporar capturas o mockups, ejecutar pruebas, validar el proceso con una tienda o publicar el servicio. Las demás secciones son el diseño y la documentación que pueden prepararse desde ahora.")
intro_heading._p.addprevious(note._p)
note.paragraph_format.space_before = Pt(5)
note.paragraph_format.space_after = Pt(10)

# Add simple verification criteria so the objectives can be checked without inventing
# savings or percentage targets.
obj3 = find_para("Objetivo específico 3:")
anchor = insert_para_after(obj3, "Criterios SMART aplicados", "Heading 3")
smart_rows = [
    ("Específico", "Indica qué funciones se completarán: inventario móvil, portal, información compartida, almacenamiento local y publicación."),
    ("Medible", "Se verificará cada operación con una prueba y un resultado observable; no se asignan porcentajes de mejora sin mediciones."),
    ("Alcanzable", "Parte del prototipo local ya construido y define como pendientes las funciones que todavía deben implementarse."),
    ("Relevante", "Responde a la necesidad de ordenar el registro y la consulta de existencias de una tienda de artículos de computación."),
    ("A tiempo", "Los objetivos se completarán y comprobarán antes de la entrega final del curso; la fecha exacta se añadirá cuando la confirme el docente."),
]
table = insert_table_after(anchor, ["Criterio", "Cómo se aplica a StockFácil PC"], smart_rows, [1.25, 5.3], 8.5)
anchor = insert_para_after(table, "Comprobación prevista de los objetivos", "Heading 3")
check_rows = [
    ("Objetivo 1", "Probar registro, búsqueda, edición y desactivación; comprobar que los datos se conservan en el teléfono sin conexión.", "Pendiente: conectar SQLite móvil y adjuntar evidencias."),
    ("Objetivo 2", "Registrar un cambio en un canal y consultar el mismo dato en el otro; repetir con productos, categorías y movimientos.", "Se reportó una prueba manual con un producto; falta adjuntar captura y documentar las demás operaciones."),
    ("Objetivo 3", "Abrir el servicio desde el teléfono y el portal fuera de la red local de desarrollo.", "Pendiente: publicar y probar el servicio y la base de datos."),
]
table2 = insert_table_after(anchor, ["Objetivo", "Criterio de comprobación", "Estado de la evidencia"], check_rows, [0.85, 3.55, 2.2], 7.8)
for cell in (table2.rows[1].cells[2], table2.rows[2].cells[2], table2.rows[3].cells[2]):
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Put the architecture in plain language within the project scope.
scope_end = find_para("La publicación del servicio y la base de datos en un alojamiento externo")
anchor = insert_para_after(scope_end, "Componentes y comunicación", "Heading 3")
table3 = insert_table_after(anchor, ["Componente", "Función y estado"], [
    ("Aplicación Android", "Permite consultar y mantener el inventario desde el teléfono cuando hay conexión."),
    ("Portal administrativo", "Permite consultar los mismos datos desde un navegador y administrar productos, categorías y movimientos."),
    ("Servicio de datos", "Recibe las solicitudes de la aplicación y del portal, valida la información y comunica los cambios."),
    ("Base central de desarrollo", "Guarda categorías, productos y movimientos en la computadora de desarrollo. Su publicación externa está pendiente."),
    ("Base local del teléfono", "Su estructura inicial existe en el código, pero no está conectada a las pantallas. El guardado sin conexión y el envío posterior están pendientes."),
], [1.65, 4.9], 8.3)
for p in table3.rows[4].cells[1].paragraphs + table3.rows[5].cells[1].paragraphs:
    for run in p.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Make unfinished work unmistakable in the sections where it affects interpretation.
for p in doc.paragraphs:
    if p.text.strip() == "Envío y procesamiento de la información de los productos.":
        p.text = "Envío y procesamiento de la información de productos, categorías y movimientos de stock."
pending_prefixes = [
    "La tienda de referencia aún no ha sido identificada",
    "El código de almacenamiento local en el teléfono",
    "La tienda usada como escenario no ha sido identificada",
    "Guardar el cambio en SQLite del teléfono",
    "Para cumplir el guardado en el teléfono",
    "La tabla SQLite del teléfono está preparada",
    "Las capturas de pantalla de cada operación",
    "La publicación del servicio y de la base de datos",
    "La matriz de pruebas con pasos",
    "La publicación del servicio y la base de datos",
    "El funcionamiento de la aplicación móvil sin conexión",
]
for prefix in pending_prefixes:
    try:
        p = find_para(prefix)
        if not p.text.strip().startswith("PENDIENTE"):
            p.text = "PENDIENTE — " + p.text
        highlight_para(p)
    except ValueError:
        pass

# The functional-requirement status cells are also yellow for the two unfinished items.
for table in doc.tables:
    if table.rows and len(table.rows[0].cells) == 3 and table.rows[0].cells[0].text.strip() == "Código":
        for row in table.rows[1:]:
            code = row.cells[0].text.strip()
            if code in {"RF-06", "RF-07"}:
                for p in row.cells[2].paragraphs:
                    for run in p.runs:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def remove_paragraphs_between(start_paragraph, stop_paragraph):
    paragraphs = doc.paragraphs
    start_index = next(i for i, p in enumerate(paragraphs) if p.text == start_paragraph.text and p.style.name == start_paragraph.style.name)
    stop_index = next(i for i, p in enumerate(paragraphs) if p.text == stop_paragraph.text and p.style.name == stop_paragraph.style.name)
    for p in paragraphs[start_index + 1:stop_index]:
        p._element.getparent().remove(p._element)


# Business use cases use the course nomenclature CUN and a single readable matrix.
cun_heading = find_para("Casos de uso de negocio")
functional_heading = find_para("6.4. Documentación funcional")
remove_paragraphs_between(cun_heading, functional_heading)
cun_rows = [
    ("CUN-01", "Mantener el catálogo\nActor: encargado de inventario", "Identifica un artículo, registra su código, nombre, categoría, marca, modelo, precio y existencias. Cuando un dato cambia, actualiza el registro.", "El artículo queda disponible para consulta."),
    ("CUN-02", "Consultar disponibilidad\nActor: encargado o responsable del portal", "Busca un artículo por nombre o código y revisa su cantidad y el nivel mínimo definido.", "El personal conoce el saldo registrado y si requiere reponer unidades."),
    ("CUN-03", "Registrar una entrada\nActor: encargado de inventario", "Cuando ingresan unidades al almacén, identifica el producto, anota la cantidad y registra una entrada con una nota opcional.", "Aumenta el saldo y queda un movimiento en el historial."),
    ("CUN-04", "Registrar una salida\nActor: encargado de inventario", "Cuando salen unidades, selecciona el producto, ingresa la cantidad y registra la salida. El proceso debe comprobar que haya saldo suficiente.", "Disminuye el saldo y queda un movimiento en el historial."),
    ("CUN-05", "Revisar movimientos\nActor: encargado o responsable del portal", "Consulta las entradas y salidas ordenadas desde la más reciente, con producto, código, cantidad, fecha y nota.", "Puede revisar los cambios registrados en las existencias."),
    ("CUN-06", "Administrar categorías\nActor: encargado o responsable del portal", "Crea, modifica o desactiva categorías. Antes de desactivar una categoría, se verifica que no tenga productos activos.", "Los productos se mantienen agrupados y consultables por categoría."),
    ("CUN-07", "Compartir información entre canales\nActor: encargado y responsable del portal", "El encargado registra un cambio desde el teléfono conectado al servicio; el responsable consulta la misma información desde el portal.", "Ambos canales muestran los datos que guardó la fuente central."),
    ("CUN-08", "Continuar sin conexión\nActor: encargado de inventario", "Guardar cambios localmente en el teléfono, marcarlos pendientes y enviarlos al servicio cuando vuelva la conexión.", "Diseño objetivo; el almacenamiento local móvil y la sincronización aún no están implementados."),
]
cun_table = insert_table_after(cun_heading, ["Código", "Caso y responsable", "Descripción del proceso", "Resultado esperado"], cun_rows, [0.58, 1.48, 3.35, 1.08], 7.7)
for p in cun_table.rows[-1].cells[3].paragraphs:
    for run in p.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

flow_title = insert_para_after(cun_table, "Diagrama de flujo de negocio: registro de entradas y salidas", "Heading 3", page_break=True)
flow_image = PATH.with_name("flujo_negocio_stockfacil.png")
make_business_flow(flow_image)
flow_para = doc.add_paragraph()
flow_para.alignment = 1
flow_para.add_run().add_picture(str(flow_image), width=Inches(5.9))
flow_title._p.addnext(flow_para._p)
caption = insert_para_after(flow_para, "Figura 6. Flujo de negocio para registrar entradas y salidas de existencias.", "Caption")
scope_note = insert_para_after(caption, "Aclaración de alcance: la aplicación no registra ventas, pagos ni comprobantes. Si la salida de unidades ocurre por una venta, el prototipo solo registra el movimiento de salida del inventario; el detalle comercial se lleva fuera de StockFácil PC.")
pending_flow = insert_para_after(scope_note, "PENDIENTE — El flujo requiere conexión con el servicio. El guardado local y el envío posterior al recuperar la conexión se incorporarán cuando se implemente la sincronización móvil.")
highlight_para(pending_flow)


# Present each detailed use case in a compact two-column specification table.
use_case_specs = {
    "CU-01. Consultar el inventario": (
        "Encargado desde Android o responsable desde el portal.",
        "El servicio está disponible y el dispositivo tiene conexión.",
        "1. El usuario abre Productos.\n2. El sistema solicita los productos activos.\n3. Muestra código, nombre, categoría, precio, cantidad y una alerta cuando la cantidad es igual o menor al mínimo.",
        "Si no hay conexión, se muestra un aviso y no se cargan datos nuevos. La versión actual no ofrece consulta sin conexión.",
        "El usuario visualiza el inventario disponible en el servicio.",
    ),
    "CU-02. Buscar un producto": (
        "Encargado de inventario o responsable del portal.",
        "Existe acceso al servicio y hay productos registrados.",
        "1. El usuario escribe parte del nombre, SKU, marca o modelo.\n2. El sistema filtra los productos coincidentes.\n3. El usuario revisa los resultados.",
        "Si no hay coincidencias, la lista queda vacía. Si falla la conexión, se informa el error.",
        "Se muestran los productos que coinciden con el texto ingresado.",
    ),
    "CU-03. Registrar un producto": (
        "Encargado desde Android o responsable desde el portal.",
        "Hay una categoría activa y conexión con el servicio.",
        "1. El usuario abre Nuevo producto.\n2. Completa categoría, SKU, nombre, marca, modelo, precio, stock y stock mínimo.\n3. El sistema valida y guarda el registro.\n4. La lista se actualiza y el otro canal puede consultar el artículo.",
        "Se rechazan campos obligatorios vacíos, valores negativos, una categoría inactiva o un SKU duplicado.",
        "El nuevo artículo queda activo y disponible para consulta.",
    ),
    "CU-04. Editar un producto": (
        "Encargado desde Android o responsable del portal.",
        "El producto está activo y existe una categoría activa.",
        "1. El usuario selecciona Editar.\n2. Modifica los datos necesarios.\n3. El sistema valida y guarda los cambios.\n4. La lista vuelve a mostrar los datos actualizados.",
        "Un SKU duplicado, datos inválidos o un producto inexistente impiden guardar. Si se cambia directamente el stock en el formulario, no se crea un movimiento de historial.",
        "Los datos del producto quedan actualizados. El ajuste directo de stock es una limitación documentada del prototipo.",
    ),
    "CU-05. Desactivar un producto": (
        "Encargado desde Android o responsable del portal.",
        "El producto está activo.",
        "1. El usuario elige Eliminar.\n2. Confirma la acción.\n3. El sistema desactiva el producto y lo oculta del catálogo.\n4. Los movimientos anteriores se conservan.",
        "Si el usuario cancela, no hay cambios. La acción no borra físicamente el registro.",
        "El producto deja de aparecer en el inventario activo y su historial se conserva.",
    ),
    "CU-06. Administrar categorías": (
        "Encargado desde Android o responsable del portal.",
        "El servicio está disponible.",
        "1. El usuario consulta las categorías.\n2. Puede crear una categoría con nombre y descripción, o editar una existente.\n3. Para desactivarla, el sistema revisa si tiene productos activos.",
        "Se rechazan nombres vacíos o duplicados. No se desactiva una categoría con productos activos.",
        "La categoría queda registrada o actualizada; si se desactiva, ya no se ofrece para nuevos productos.",
    ),
    "CU-07. Registrar un movimiento de stock": (
        "Encargado desde Android o responsable del portal.",
        "El producto está activo y el servicio está disponible.",
        "1. El usuario selecciona producto y tipo de movimiento.\n2. Indica cantidad y una nota opcional.\n3. El sistema valida el tipo y que la cantidad sea mayor que cero.\n4. Para una salida, comprueba que haya saldo suficiente.\n5. Guarda el movimiento y actualiza las existencias.",
        "No se acepta un tipo distinto de entrada o salida, una cantidad menor que uno, un producto inactivo o una salida superior al saldo.",
        "El movimiento aparece en el historial y el saldo del producto se actualiza.",
    ),
    "CU-08. Consultar movimientos": (
        "Encargado de inventario o responsable del portal.",
        "El servicio está disponible.",
        "1. El usuario abre Movimientos.\n2. El sistema muestra los registros desde el más reciente, con producto, SKU, tipo, cantidad, fecha y nota.",
        "Si no hay registros, se indica que el historial está vacío. La interfaz no permite editar ni eliminar movimientos.",
        "El usuario revisa los cambios de existencias que se registraron.",
    ),
    "CU-09. Sincronizar cambios guardados en el teléfono": (
        "Encargado de inventario.",
        "El almacenamiento local y la sincronización están implementados y vuelve la conexión.",
        "1. La aplicación identifica cambios locales pendientes.\n2. Envía cada cambio al servicio.\n3. El servicio valida y guarda la información.\n4. La aplicación marca como enviados los cambios confirmados y actualiza la vista.",
        "La forma de resolver cambios simultáneos sobre un mismo producto debe definirse antes de cerrar la sincronización.",
        "Diseño previsto; no disponible en el prototipo actual. Requiere integrar SQLite con las pantallas y probar la sincronización.",
    ),
}
for heading in [p for p in doc.paragraphs if p.style.name == "Heading 3" and p.text.strip().startswith("CU-")]:
    title = next((name for name in use_case_specs if heading.text.strip().startswith(name)), None)
    if title is None:
        continue
    if title.startswith("CU-05"):
        heading.paragraph_format.page_break_before = True
    current = doc.paragraphs
    i = next(i for i, p in enumerate(current) if p.text == heading.text and p.style.name == heading.style.name)
    remove = []
    for p in current[i + 1:]:
        if p.style.name.startswith("Heading"):
            break
        if p.text.strip():
            remove.append(p)
    for p in remove:
        p._element.getparent().remove(p._element)
    actor, precondition, flow, alternative, result = use_case_specs[title]
    table = insert_table_after(heading, ["Elemento", "Especificación"], [
        ("Actor", actor),
        ("Precondición", precondition),
        ("Flujo principal", flow),
        ("Alternativas", alternative),
        ("Resultado / estado", result),
    ], [1.25, 5.25], 8.1)
    if title.startswith("CU-09"):
        for p in table.rows[-1].cells[1].paragraphs:
            for run in p.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


# Rules are presented with their meaning and current implementation status.
rules_heading = find_para("6.7. Reglas de negocio")
data_heading = find_para("6.8. Modelo lógico y físico de datos")
remove_paragraphs_between(rules_heading, data_heading)
rule_rows = [
    ("RN-01", "Cada producto pertenece a una categoría activa.", "Se valida al registrar o actualizar el producto. Implementada."),
    ("RN-02", "El SKU es obligatorio y no se repite.", "El servicio comprueba que el código no exista en otro producto. Implementada."),
    ("RN-03", "El nombre de la categoría es obligatorio y único.", "El servicio y la base de datos impiden nombres duplicados. Implementada."),
    ("RN-04", "Precio, existencias y stock mínimo no pueden ser negativos.", "Los valores se validan antes de guardar. Implementada."),
    ("RN-05", "Hay stock bajo cuando la cantidad es igual o menor al mínimo definido.", "El resumen identifica los productos que requieren atención. Implementada."),
    ("RN-06", "Un movimiento requiere un producto activo, un tipo válido y una cantidad mayor que cero.", "Solo se aceptan entradas y salidas. Implementada."),
    ("RN-07", "Una salida no puede superar las existencias disponibles.", "El sistema rechaza la operación si el saldo no alcanza. Implementada."),
    ("RN-08", "Una entrada o salida aceptada guarda el movimiento y actualiza las existencias.", "El historial registra tipo, cantidad, fecha y nota. Implementada."),
    ("RN-09", "No se desactiva una categoría que tenga productos activos.", "El sistema verifica los productos relacionados. Implementada."),
    ("RN-10", "Desactivar un producto conserva el registro y sus movimientos anteriores.", "El producto se oculta del catálogo activo sin borrar su historial. Implementada."),
    ("RN-11", "El historial se consulta desde el movimiento más reciente.", "La interfaz permite consultar, pero no editar ni borrar movimientos. Implementada con ese límite."),
    ("RN-12", "El historial debe permitir distinguir movimientos de cambios directos al stock.", "Hoy el formulario permite cambiar existencias sin crear un movimiento. Limitación: el historial no refleja esos ajustes directos."),
    ("RN-13", "El teléfono y el portal usan datos compartidos; el trabajo sin conexión requiere guardar y enviar cambios al recuperar la red.", "La versión actual requiere conexión. El guardado local y la sincronización están pendientes."),
]
rules_table = insert_table_after(rules_heading, ["Código", "Regla de negocio", "Aplicación y estado"], rule_rows, [0.62, 2.75, 3.13], 7.8)
for idx in (12, 13):
    for cell in rules_table.rows[idx].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


# Clarify where the documentation, diagrams, and screenshots belong in the report.
tech_intro = find_para("La documentación del proyecto reúne el diseño funcional")
tech_note = insert_para_after(tech_intro, "Ubicación de los entregables: los modelos lógico y físico y el diagrama entidad-relación están desarrollados en la sección 6.8; los casos de uso y reglas aparecen en las secciones 6.3, 6.5 y 6.7. Las capturas reales, los mockups y la evidencia de pruebas se ubicarán en los anexos y se mencionarán desde la sección de productos y entregables. La sección 7.6 resume los documentos entregados; no reemplaza los modelos ni las evidencias.")

# Append glossary, bibliography with usable links, and annexes as a completion checklist.
last_recommendation = find_para("Publicar el servicio y la base de datos en un alojamiento aprobado")
anchor = insert_para_after(last_recommendation, "10. GLOSARIO", "Heading 1")
glossary_rows = [
    ("API / servicio de datos", "Parte del sistema que recibe solicitudes de la aplicación y del portal, y les entrega o guarda información."),
    ("REST", "Forma habitual de organizar solicitudes entre programas a través de una red."),
    ("SQLite", "Base de datos ligera que guarda información en un archivo. En este proyecto se usa en la computadora para el prototipo y está prevista también para el teléfono."),
    ("SKU", "Código único que identifica cada artículo del inventario."),
    ("Stock mínimo", "Cantidad de referencia que permite advertir que quedan pocas unidades de un producto."),
    ("Movimiento de stock", "Registro de una entrada o salida de unidades de un producto."),
    ("Desactivación lógica", "Acción que oculta un producto del catálogo sin borrar su información ni su historial."),
    ("Sincronización", "Envío y actualización de los cambios entre el teléfono y la información compartida cuando vuelve la conexión."),
    ("Modelo entidad-relación", "Dibujo que muestra qué datos guarda el sistema y cómo se relacionan entre sí."),
]
table4 = insert_table_after(anchor, ["Término", "Explicación sencilla"], glossary_rows, [1.8, 4.75], 8.4)
anchor = insert_para_after(table4, "11. BIBLIOGRAFÍA", "Heading 1", page_break=True)
sources = [
    ("Instituto Nacional de Estadística e Informática (2026). El 64,5 % de los hogares del país tiene conexión a Internet durante el segundo trimestre de 2026. Resultados de la Encuesta Nacional de Hogares.", "https://www.gob.pe/institucion/inei/noticias/1450362-el-64-5-de-los-hogares-del-pais-tiene-conexion-a-internet-durante-el-segundo-trimestre-de-2026"),
    ("Ministerio de la Producción (2025a). Número de Mipyme en Perú ascendió a 2,33 millones al cierre de 2024.", "https://www.gob.pe/institucion/produce/noticias/1299194-numero-de-mipyme-en-peru-ascendio-a-2-33-millones-al-cierre-del-2024-alcanzando-un-crecimiento-de-1-6"),
    ("Ministerio de la Producción, OGEIEE (2025b). Las Mipyme en Cifras 2024. Primera edición, octubre de 2025; resultados de la EEA 2024 correspondientes al ejercicio 2023.", "https://www.producempresarial.pe/wp-content/uploads/2025/11/La-Mipyme-en-Cifras-2024_.pdf"),
    ("Presidencia del Consejo de Ministros (2023). Decreto Supremo N.° 085-2023-PCM, Política Nacional de Transformación Digital al 2030.", "https://www.gob.pe/institucion/pcm/normas-legales/4471543-085-2023-pcm"),
    ("Autoridad Nacional de Protección de Datos Personales (2025). Nuevo Reglamento de Protección de Datos Personales; vigencia desde el 31 de marzo de 2025.", "https://www.gob.pe/institucion/anpd/campa%C3%B1as/128319-nuevo-reglamento-de-proteccion-de-datos-personales"),
    ("Ministerio de Justicia y Derechos Humanos (2024). Decreto Supremo N.° 016-2024-JUS, Reglamento de la Ley de Protección de Datos Personales.", "https://www.gob.pe/institucion/anpd/normas-legales/6554453-n-016-2024-jus"),
    ("Ministerio del Ambiente (2025). Minam promueve el reciclaje de residuos de aparatos eléctricos y electrónicos; cifras de Sigersol de 2020 a 2024.", "https://www.gob.pe/institucion/minam/noticias/1263803-minam-promueve-reciclaje-de-residuos-de-aparatos-electricos-y-electronicos"),
    ("Ministerio del Ambiente (2019). Decreto Supremo N.° 009-2019-MINAM, Régimen Especial de Gestión y Manejo de Residuos de Aparatos Eléctricos y Electrónicos.", "https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam"),
    ("Liu, F., Yu, Y., Fang, Y., Zhu, M., Shi, Y. y Xiao, S. (2024). Lean strategy in SMEs: Inventory leanness, operational leanness, and financial performance. The Asian Journal of Shipping and Logistics, 40(2), 109–117.", "https://doi.org/10.1016/j.ajsl.2024.02.003"),
]
for citation, url in sources:
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run(citation + " ")
    add_hyperlink(p, "Abrir fuente oficial", url)
    p.paragraph_format.keep_together = True
    anchor._p.addnext(p._p)
    anchor = p

anchor = insert_para_after(anchor, "Cibertec (2026). Desarrollo de Aplicaciones Móviles I: guía del proyecto y rúbrica del curso. Material proporcionado por el estudiante.", "List Paragraph")
anchor = insert_para_after(anchor, "12. ANEXOS", "Heading 1", page_break=True)
anchor = insert_para_after(anchor, "Los anexos siguientes se completarán con evidencia real de la versión que se implemente. No se han creado capturas ficticias.")

annexes = [
    ("Anexo A. Mockups de la aplicación y del portal", "PENDIENTE — Elaborar y adjuntar las pantallas acordadas en la herramienta de diseño. Las imágenes se incorporarán cuando se revisen juntos."),
    ("Anexo B. Evidencias de la aplicación Android", "PENDIENTE — Adjuntar capturas reales de inicio/resumen, lista y búsqueda de productos, formulario de producto, categorías y movimientos, indicando el dispositivo usado."),
    ("Anexo C. Evidencias del portal y de los datos compartidos", "PENDIENTE — Adjuntar capturas del portal y una secuencia que muestre el mismo producto después de registrarlo desde el teléfono. La prueba manual ya fue reportada; falta insertar su evidencia."),
    ("Anexo D. Evidencia de SQLite local en el teléfono", "PENDIENTE — Completar la integración de SQLite con las pantallas móviles y adjuntar una comprobación de guardado y consulta sin conexión. La tabla que existe en el código no demuestra por sí sola esta función."),
    ("Anexo E. Matriz de pruebas", "La matriz define escenarios y resultados esperados. Las columnas de resultado real y evidencia se completarán después de ejecutar cada prueba.")
]
for title, text in annexes:
    anchor = insert_para_after(anchor, title, "Heading 2")
    p = insert_para_after(anchor, text)
    if text.startswith("PENDIENTE") or "se completarán" in text:
        highlight_para(p)
    anchor = p

test_rows = [
    ("PT-01", "Registrar producto con datos válidos", "El producto aparece en la lista y el portal puede consultarlo.", "Pendiente de ejecutar y adjuntar captura."),
    ("PT-02", "Intentar guardar un SKU repetido", "El sistema rechaza el duplicado y mantiene el registro anterior.", "Pendiente de ejecutar y adjuntar captura."),
    ("PT-03", "Registrar una entrada de stock", "Aumenta la existencia y aparece el movimiento en el historial.", "Pendiente de ejecutar y adjuntar captura."),
    ("PT-04", "Registrar una salida mayor que el stock disponible", "El sistema rechaza la salida y conserva la cantidad previa.", "Pendiente de ejecutar y adjuntar captura."),
    ("PT-05", "Desactivar un producto con movimientos previos", "El producto deja de aparecer en el catálogo y conserva el historial.", "Pendiente de ejecutar y adjuntar captura."),
    ("PT-06", "Guardar un cambio sin conexión y recuperar la conexión", "El teléfono conserva el cambio y lo envía al servicio al reconectarse.", "No ejecutable todavía: falta integrar el almacenamiento local y la sincronización."),
]
test_table = insert_table_after(anchor, ["ID", "Escenario", "Resultado esperado", "Resultado real / evidencia"], test_rows, [0.5, 1.7, 2.6, 1.7], 7.5)
for row in test_table.rows[1:]:
    for p in row.cells[3].paragraphs:
        for run in p.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
anchor = insert_para_after(test_table, "Anexo F. Video de demostración", "Heading 2")
p = insert_para_after(anchor, "PENDIENTE — Preparar el video de demostración de 3 a 5 minutos indicado por la guía del curso. Debe mostrar las funciones que estén implementadas y probadas al momento de grabarlo.")
highlight_para(p)
anchor = p
anchor = insert_para_after(anchor, "Anexo G. Validación del proceso con una tienda", "Heading 2")
p = insert_para_after(anchor, "PENDIENTE — Identificar y consultar una tienda participante; registrar cómo controla hoy su inventario y qué necesita comprobar. Con esa información se podrán validar el diagnóstico y los beneficios esperados sin atribuirle cifras no medidas.")
highlight_para(p)

# Remove trailing empty paragraphs that created a blank final page in the source draft.
for p in reversed(doc.paragraphs):
    if p.text.strip():
        break
    p._element.getparent().remove(p._element)

# Apply consistent spacing and protect headings from being orphaned.
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True
    if p.text.strip().startswith("9. RECOMENDACIONES"):
        p.paragraph_format.page_break_before = True
    if p.style.name == "Normal":
        p.paragraph_format.widow_control = True

# Avoid stretched word spacing in any table imported from the original draft.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.keep_together = True
                p.paragraph_format.widow_control = True

doc.core_properties.title = "Informe del proyecto StockFácil PC — borrador de trabajo"
doc.core_properties.subject = "Diseño funcional, estado del prototipo y pendientes de evidencia"
doc.save(PATH)
flow_image.unlink(missing_ok=True)
print(PATH)
print(f"Párrafos: {len(doc.paragraphs)} | Tablas: {len(doc.tables)} | Figuras: {len(doc.inline_shapes)}")
