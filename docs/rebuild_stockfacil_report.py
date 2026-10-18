from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(r"C:\Users\Usuario\Downloads\Informe_Proyecto_StockFacil_PC_Avance_SEPTE (1).docx")
OUT = Path(r"C:\Users\Usuario\Documents\Codex\2026-10-04\hola-me-gustar-a-saber-si\outputs\StockFacil\docs\Informe_StockFacil_PC_Avance_con_documentacion_funcional.docx")
ER_IMAGE = OUT.with_name("modelo_er_stockfacil.png")
SOCIAL_CHART = OUT.with_name("_septe_social_chart.png")
RAEE_CHART = OUT.with_name("_septe_raee_chart.png")


def find_paragraph(doc, starts):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(starts):
            return paragraph
    raise ValueError(f"No se encontró el párrafo: {starts}")


def replace_paragraph(doc, starts, text):
    paragraph = find_paragraph(doc, starts)
    paragraph.text = text
    return paragraph


def insert_paragraph(doc, anchor, text="", style="Normal", bold_prefix=None, page_break=False):
    paragraph = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    elif text:
        paragraph.add_run(text)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    anchor_element = anchor._p if hasattr(anchor, "_p") else anchor._tbl
    anchor_element.addnext(paragraph._p)
    return paragraph


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def insert_table(doc, anchor, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255), size=font_size)
        shade_cell(table.rows[0].cells[index], "23466F")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value), size=font_size)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    anchor._p.addnext(table._tbl)
    return table


def create_er_image(path):
    width, height = 1800, 650
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    regular = ImageFont.truetype(regular_path, 24)
    small = ImageFont.truetype(regular_path, 21)
    heading = ImageFont.truetype(bold_path, 27)
    title = ImageFont.truetype(bold_path, 34)
    draw.text((width // 2, 24), "MODELO ENTIDAD-RELACIÓN", font=title, fill="#193653", anchor="mt")

    def entity(x, y, w, h, name, fields):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill="#F7FAFD", outline="#23466F", width=4)
        draw.rounded_rectangle((x, y, x + w, y + 65), radius=16, fill="#23466F")
        draw.rectangle((x, y + 48, x + w, y + 65), fill="#23466F")
        draw.text((x + w / 2, y + 14), name, font=heading, fill="white", anchor="mt")
        cy = y + 86
        for field in fields:
            draw.text((x + 24, cy), field, font=small, fill="#1F2937")
            cy += 38

    entity(45, 125, 490, 410, "CATEGORÍA", [
        "PK  Id", "Nombre", "Descripción", "Activa"
    ])
    entity(655, 125, 490, 410, "PRODUCTO", [
        "PK  Id", "FK  CategoryId", "SKU, nombre, marca, modelo", "Precio, stock, stock mínimo", "Activo"
    ])
    entity(1265, 125, 490, 410, "MOVIMIENTO DE STOCK", [
        "PK  Id", "FK  ProductId", "Tipo (IN / OUT)", "Cantidad y fecha", "Nota"
    ])
    y = 330
    draw.line((535, y, 655, y), fill="#345F87", width=5)
    draw.polygon([(655, y), (638, y - 11), (638, y + 11)], fill="#345F87")
    draw.text((545, y - 45), "1", font=regular, fill="#193653")
    draw.text((625, y - 45), "N", font=regular, fill="#193653")
    draw.line((1145, y, 1265, y), fill="#345F87", width=5)
    draw.polygon([(1265, y), (1248, y - 11), (1248, y + 11)], fill="#345F87")
    draw.text((1155, y - 45), "1", font=regular, fill="#193653")
    draw.text((1235, y - 45), "N", font=regular, fill="#193653")
    draw.text((width // 2, 585), "Una categoría agrupa productos; cada producto puede tener varios movimientos registrados.", font=regular, fill="#46586B", anchor="mt")
    image.save(path, "PNG")


def replace_embedded_image(doc, paragraph_prefix, image_path):
    found_anchor = False
    blips = []
    for paragraph in doc.paragraphs:
        if not found_anchor and paragraph.text.strip().startswith(paragraph_prefix):
            found_anchor = True
            continue
        if found_anchor:
            blips = paragraph._p.xpath(".//a:blip")
            if blips:
                break
    if not blips:
        raise ValueError(f"No se encontró una imagen asociada a: {paragraph_prefix}")
    rel_id = blips[0].get(qn("r:embed"))
    doc.part.related_parts[rel_id]._blob = Path(image_path).read_bytes()


def create_horizontal_bar_chart(path, title, entries, max_value, color):
    width, height = 988, 574
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 20)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 25)
    small = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 17)
    draw.text((width // 2, 24), title, font=bold, fill="#17365D", anchor="mt")

    left, right, top, bottom = 280, 64, 100, 92
    plot_width = width - left - right
    plot_height = height - top - bottom
    row_height = plot_height / len(entries)
    for tick in range(0, max_value + 1, 20):
        x = left + plot_width * tick / max_value
        draw.line((x, top - 4, x, height - bottom), fill="#D6DEE8", width=2)
        draw.text((x, height - bottom + 12), str(tick), font=small, fill="#46586B", anchor="mt")

    for index, (label, value) in enumerate(entries):
        y_mid = top + row_height * (index + 0.5)
        bar_height = 52
        bar_right = left + plot_width * value / max_value
        draw.text((left - 16, y_mid), label, font=regular, fill="#27384A", anchor="rm")
        draw.rounded_rectangle((left, y_mid - bar_height / 2, bar_right, y_mid + bar_height / 2), radius=8, fill=color)
        draw.text((bar_right + 10, y_mid), f"{value:.1f} %".replace(".", ","), font=regular, fill="#27384A", anchor="lm")

    draw.text((left + plot_width / 2, height - 27), "Porcentaje de hogares", font=small, fill="#46586B", anchor="mm")
    image.save(path, "PNG")


def create_raee_chart(path):
    width, height = 988, 574
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 19)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 25)
    small = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 16)
    draw.text((width // 2, 24), "RAEE reciclados en el Perú, 2020–2024", font=bold, fill="#17365D", anchor="mt")

    left, right, top, bottom = 92, 42, 94, 88
    plot_width = width - left - right
    plot_height = height - top - bottom
    max_value = 30000
    for tick in range(0, max_value + 1, 5000):
        y = top + plot_height * (1 - tick / max_value)
        draw.line((left, y, width - right, y), fill="#D6DEE8", width=2)
        draw.text((left - 12, y), f"{tick:,}".replace(",", " "), font=small, fill="#46586B", anchor="rm")

    values = [5700, 11175, 14700, 20000, 25600]
    labels = ["5 700", "11 175", ">14 700", ">20 000", ">25 600"]
    years = ["2020", "2021", "2022", "2023", "2024"]
    bar_width = 92
    gap = (plot_width - len(values) * bar_width) / (len(values) + 1)
    for index, (value, label, year) in enumerate(zip(values, labels, years)):
        x0 = left + gap * (index + 1) + bar_width * index
        x1 = x0 + bar_width
        y1 = top + plot_height
        y0 = top + plot_height * (1 - value / max_value)
        draw.rounded_rectangle((x0, y0, x1, y1), radius=7, fill="#2F7D63")
        draw.text(((x0 + x1) / 2, y0 - 12), label, font=small, fill="#27384A", anchor="ms")
        draw.text(((x0 + x1) / 2, y1 + 18), year, font=regular, fill="#27384A", anchor="mt")
    draw.text((left + 5, top - 22), "Toneladas", font=small, fill="#46586B", anchor="lm")
    image.save(path, "PNG")


doc = Document(SOURCE)

create_horizontal_bar_chart(
    SOCIAL_CHART,
    "Hogares con Internet por zona (II trimestre de 2026)",
    [("Lima Metropolitana", 84.7), ("Resto urbano", 66.2), ("Área rural", 27.1)],
    100,
    "#315F91",
)
create_raee_chart(RAEE_CHART)
replace_embedded_image(doc, "El acceso a Internet en los hogares peruanos", SOCIAL_CHART)
replace_embedded_image(doc, "Los artículos de computación forman parte", RAEE_CHART)

# Clarify project status in the narrative while preserving the original SEPTE figures and citations.
replace_paragraph(doc, "Este informe presenta StockFácil PC", "Este informe presenta StockFácil PC, un proyecto aplicado que propone ordenar el control del inventario de una pequeña tienda de artículos de computación. El diagnóstico revisa factores sociales, económicos, políticos, tecnológicos y ecológicos con información de fuentes oficiales. Estos datos ayudan a entender el entorno nacional, pero no reemplazan la validación de las necesidades de una tienda específica.")
replace_paragraph(doc, "A partir de ese diagnóstico", "A partir del diagnóstico, se plantean objetivos SMART —específicos, medibles, alcanzables, relevantes y con un plazo definido— para organizar el registro y la consulta de productos y existencias. La propuesta busca facilitar la búsqueda de información, mantener al día las entradas y salidas y ayudar a tomar decisiones sobre los productos disponibles. La reducción de demoras o errores es un beneficio esperado, todavía no medido en un negocio participante.")
replace_paragraph(doc, "El avance reúne una aplicación móvil", "El avance reúne una aplicación móvil y un portal administrativo conectados por un servicio de datos. En las pruebas realizadas en la red local, un producto registrado desde el teléfono se pudo consultar en el portal. La versión actual guarda los datos compartidos en una base de desarrollo ubicada en la computadora. Aunque el proyecto Android contiene una tabla SQLite preparada para productos, la pantalla móvil todavía no la utiliza; por eso el guardado sin conexión y la sincronización posterior siguen pendientes. El informe desarrolla los objetivos, la justificación, el alcance, los casos de uso, los flujos, las reglas de negocio y el modelo de datos, y separa esos planos de las funciones ya construidas.")
replace_paragraph(doc, "StockFácil PC reúne una aplicación móvil", "StockFácil PC reúne una aplicación móvil para consultar y mantener el inventario, un portal web para administrar la misma información y un servicio que comunica ambos canales. En la versión actual, el teléfono consulta y envía cambios al servicio cuando hay conexión; el servicio guarda los datos compartidos en la computadora y el portal los consulta desde allí. El proyecto también contempla guardar información en el teléfono y enviarla al recuperar la conexión, función que todavía debe construirse.")
replace_paragraph(doc, "El informe completo también documentará", "Este informe presenta el diagnóstico, los objetivos, la justificación y los planos funcionales y de datos que orientan el desarrollo. Los casos de uso y flujos describen el comportamiento esperado del sistema; no son, por sí mismos, evidencia de que cada función esté terminada. La tienda de referencia aún no ha sido identificada ni entrevistada, por lo que el diagnóstico no atribuye a una empresa concreta pérdidas, errores o tiempos que no se hayan medido. Los mockups y las capturas de evidencia se incorporarán en una etapa posterior.")

# Expand the SEPTE diagnosis with newer, scoped evidence and clear local limits.
replace_paragraph(doc, "El acceso a Internet en los hogares peruanos", "Según la Encuesta Nacional de Hogares citada por el INEI (2026), en el segundo trimestre de 2026 el 64,5 % de los hogares peruanos tenía acceso a Internet. La cifra fue de 84,7 % en Lima Metropolitana, 66,2 % en el resto urbano y 27,1 % en el área rural. Además, el 90,2 % de las personas usuarias de Internet de seis años a más accedió desde un teléfono celular. La diferencia por zona hace pertinente considerar el teléfono como canal de consulta; estas cifras nacionales no describen la conectividad de una tienda específica.")
replace_paragraph(doc, "Figura 1.", "Figura 1. Hogares con acceso a Internet por zona de residencia, segundo trimestre de 2026.")
replace_paragraph(doc, "Fuente: INEI, Estadísticas de las Tecnologías", "Fuente: INEI (2026), resultados de la Encuesta Nacional de Hogares. Elaboración propia a partir del informe de tecnologías de información y comunicación en los hogares.")
replace_paragraph(doc, "La diferencia entre zonas muestra", "La brecha por zona muestra que el acceso no es uniforme. La alta proporción de usuarios que emplea celular permite justificar que el proyecto considere ese canal, pero no basta para afirmar que la tienda cuenta con conexión estable. El prototipo actual requiere que el teléfono y la computadora estén conectados al servicio local; el trabajo sin conexión todavía está pendiente.")
replace_paragraph(doc, "El Ministerio de la Producción analizó", "El Ministerio de la Producción reportó que en 2024 operaron 2 331 173 Mipyme formales, equivalentes al 99,3 % de las empresas formales del país (PRODUCE, 2025a). Por otro lado, la Encuesta Económica Anual 2024, cuyos resultados corresponden al ejercicio 2023, consultó a Mipyme que no utilizaban plataformas digitales para vender: el 37,7 % no percibía mayores beneficios, el 21,4 % consideraba elevado el precio de adquisición y el 12,4 % mencionó los costos de operación y mantenimiento (PRODUCE, 2025b). La pregunta permitía indicar más de un motivo.")
replace_paragraph(doc, "Fuente: PRODUCE, Las Mipyme en Cifras", "Fuente: PRODUCE–OGEIEE, Las Mipyme en Cifras 2024, gráfico 8.13; resultados de la EEA 2024 (ejercicio 2023). Pregunta dirigida a Mipyme no usuarias de plataformas digitales de venta; respuesta múltiple. Elaboración propia.")
replace_paragraph(doc, "El resultado indica que el valor percibido", "Las cifras describen motivos para no utilizar plataformas de venta digital; no miden dificultades del control interno de inventario ni el costo de una aplicación como StockFácil PC. Sirven como contexto para tomar en cuenta el costo y el beneficio percibido al proponer herramientas para Mipyme. El ahorro del proyecto solo podrá estimarse después de conocer cuánto tiempo y dinero dedica una tienda a registrar, buscar y corregir sus datos, y comparar ese punto de partida con el uso de la solución.")
replace_paragraph(doc, "El marco público peruano promueve", "En 2023 se aprobó la Política Nacional de Transformación Digital al 2030 mediante el Decreto Supremo N.° 085-2023-PCM. La política orienta la transformación digital en el país y considera a distintos sectores, incluido el privado, en cuanto les sea aplicable (PCM, 2023). Se presenta aquí como contexto nacional, no como una exigencia de que este proyecto incorpore funciones comerciales o administrativas ajenas a su alcance.")
replace_paragraph(doc, "La Política Nacional de Transformación Digital", "El nuevo Reglamento de la Ley de Protección de Datos Personales, aprobado por el Decreto Supremo N.° 016-2024-JUS, entró en vigencia el 31 de marzo de 2025 (ANPD, 2025). El prototipo de StockFácil PC registra categorías, productos y movimientos de inventario; no incluye datos personales de clientes o trabajadores, cuentas ni inicio de sesión. Esta descripción delimita la información manejada actualmente y no atribuye al prototipo funciones de gestión de usuarios.")
replace_paragraph(doc, "La Encuesta Económica Anual 2024", "La Encuesta Económica Anual 2024, publicada por PRODUCE en 2025, reportó para el ejercicio 2023 que el 90,1 % de las Mipyme contaba con Internet, el 85,1 % con teléfono móvil, el 70,9 % con computadora de escritorio y el 60,3 % con laptop. En el mismo informe, el 6,6 % indicó haber realizado ventas por Internet y el 6,4 % compras por ese medio (PRODUCE, 2025b). Son indicadores de acceso y comercio electrónico; no miden si las empresas llevan su inventario en un sistema digital.")
replace_paragraph(doc, "La presencia de teléfonos y computadoras", "El acceso de las Mipyme a teléfonos y computadoras ofrece contexto para diseñar soluciones que puedan usarse desde distintos equipos. Sin embargo, contar con dispositivos o Internet no significa que el inventario ya se gestione de forma digital: la encuesta no mide ese proceso. En StockFácil PC, el teléfono y el portal comparten información cuando se conectan al servicio en la red local; la tabla SQLite del teléfono aún no está conectada a las pantallas y no se presenta como almacenamiento móvil funcional ni como evidencia de uso sin conexión.")
replace_paragraph(doc, "Los artículos de computación forman parte", "Los productos electrónicos requieren una gestión adecuada cuando se convierten en residuos. El Ministerio del Ambiente informa que en el Perú se generan más de 100 mil toneladas de RAEE al año. Según los registros de Sigersol citados por MINAM (2025), se reciclaron 5 700 toneladas en 2020, 11 175 en 2021, más de 14 700 en 2022, más de 20 000 en 2023 y más de 25 600 en 2024. El Decreto Supremo N.° 009-2019-MINAM regula el manejo especial de estos residuos y asigna responsabilidades a los actores correspondientes de esa cadena.")
replace_paragraph(doc, "Figura 4.", "Figura 4. RAEE reciclados en el Perú, 2020–2024.")
replace_paragraph(doc, "Fuente: MINAM, registros del Sigersol.", "Fuente: MINAM (2025), cifras de Sigersol. Para 2022, 2023 y 2024, la entidad publica cantidades mínimas expresadas como “más de”; las barras muestran esos mínimos y no valores exactos. Elaboración propia.")
replace_paragraph(doc, "Este contexto es pertinente para un negocio", "Estas cifras muestran la importancia de manejar adecuadamente los residuos electrónicos, pero no describen la situación de una tienda participante ni prueban que el sistema reduzca esos residuos. StockFácil PC registra productos y movimientos de inventario; no registra residuos ni su entrega a operadores autorizados. Si se reemplazaran anotaciones en papel, el posible efecto ambiental tendría que medirse observando el proceso real.")
replace_paragraph(doc, "La información oficial muestra", "Las fuentes consultadas muestran que el celular es un medio frecuente de acceso a Internet, que las Mipyme tienen presencia importante en la economía formal y que el acceso a dispositivos no equivale por sí solo a digitalizar el control de inventario. También describen normas de transformación digital y protección de datos, además de un aumento de las cantidades de RAEE recicladas. Estos datos justifican estudiar una solución móvil y un portal administrativo, pero no demuestran una falla específica en una tienda ni permiten prometer ahorros. Para completar el diagnóstico local se deberá observar cómo se registra el inventario y medir el tiempo y los errores del proceso actual.")

# Keep the political/legal table aligned with the implemented data scope.
for table in doc.tables:
    if table.rows and any(cell.text.strip() == "Marco público" for cell in table.rows[0].cells):
        if len(table.rows) >= 3 and len(table.columns) >= 3:
            table.cell(1, 2).text = "Orienta la transformación digital como contexto nacional; no define una función específica del prototipo."
            table.cell(2, 2).text = "El modelo actual no incluye datos personales de clientes o trabajadores; registra información de inventario."
        break
replace_paragraph(doc, "Estos instrumentos brindan contexto", "La Política Nacional de Transformación Digital y el reglamento de protección de datos forman parte del contexto que guía el diseño. El alcance de datos implementado en el prototipo comprende categorías, productos y movimientos de inventario; actualmente no incluye cuentas de clientes o trabajadores ni acceso con inicio de sesión. Por esa razón, este informe no afirma que exista un módulo de usuarios ni que se haya completado una evaluación legal o de seguridad para publicar el servicio.")
replace_paragraph(doc, "La presencia de teléfonos y computadoras", "La presencia de teléfonos y computadoras en las Mipyme respalda una solución que pueda consultarse desde más de un equipo. En el prototipo, el teléfono y el portal comparten información cuando se conectan a la API en la red local. El código de SQLite del teléfono aún no está conectado a las pantallas, por lo que no se presenta como almacenamiento móvil funcional ni se afirma que la aplicación opere sin conexión.")
replace_paragraph(doc, "El desarrollo de StockFácil PC surge", "StockFácil PC se propone para facilitar el registro y la consulta de productos y existencias en una tienda de artículos de computación. Al reunir la información del catálogo y los movimientos en un mismo sistema, la solución puede ayudar a que el personal encuentre los datos con mayor facilidad. Este aporte se plantea como propósito del proyecto; su efecto en tiempos, errores o costos deberá medirse cuando se valide con una tienda.")
replace_paragraph(doc, "La implementación de StockFácil PC se justifica", "La implementación de StockFácil PC se justifica por la necesidad de consultar y mantener la información del inventario desde el teléfono y desde una computadora. En el avance actual ambos canales acceden a los mismos datos cuando están conectados a la red local, lo que permite revisar los cambios desde el portal después de registrarlos en la aplicación móvil.")
justification_paragraph = replace_paragraph(doc, "Asimismo, el uso de SQLite en la aplicación móvil", "La propuesta reúne los datos de productos, categorías y movimientos para que el personal consulte una referencia común. El almacenamiento local del teléfono y el envío automático de cambios al recuperar la conexión forman parte del diseño previsto, pero aún no están integrados en el flujo móvil actual. Por ello, no se atribuyen todavía ahorros ni mejoras cuantificadas al sistema.")
insert_paragraph(doc, justification_paragraph, "Como antecedente, Liu y colaboradores (2024) analizaron 4 019 pequeñas y medianas empresas manufactureras en China. Encontraron que mantener menos unidades puede relacionarse con mejores resultados hasta cierto punto, pero reducir demasiado las existencias puede perjudicarlos. El estudio corresponde a otro país y rubro, por lo que no permite anticipar cuánto mejorará una tienda peruana. Para StockFácil PC, su aporte es una idea práctica: mostrar saldos y movimientos para que el encargado decida con información, sin asumir que tener menos stock siempre es mejor.")
replace_paragraph(doc, "Finalmente, StockFácil PC representa", "El proyecto permite aplicar los conocimientos del curso en el desarrollo de una solución de inventario que conecta una aplicación Android con un portal administrativo. La justificación se centra en el beneficio que se busca para el negocio: contar con información ordenada y disponible desde ambos canales, sujeto a completar las funciones pendientes y validar el proceso con usuarios de una tienda.")

# Objectives are goals for the full project, not claims that all work is already complete.
replace_paragraph(doc, "Desarrollar el sistema StockFácil PC", "Desarrollar e integrar StockFácil PC para que una tienda de artículos de computación pueda registrar, consultar y mantener actualizados sus productos y existencias desde una aplicación móvil y un portal administrativo, con información compartida y almacenamiento local en el teléfono antes de la entrega final del curso.")
replace_paragraph(doc, "Objetivo específico 1:", "Objetivo específico 1: Implementar en la aplicación móvil el registro, búsqueda, edición y eliminación lógica de productos, junto con su almacenamiento local en SQLite, y comprobar cada operación mediante casos funcionales antes de la entrega final del proyecto.")
replace_paragraph(doc, "Objetivo específico 2:", "Objetivo específico 2: Integrar la aplicación móvil con el portal administrativo para productos, categorías y movimientos de stock, y comprobar antes de la entrega que los cambios enviados desde el teléfono se puedan consultar en el portal.")
replace_paragraph(doc, "Objetivo específico 3:", "Objetivo específico 3: Publicar la base de datos y el servicio de integración en un alojamiento gratuito compatible con el proyecto, y comprobar su acceso desde el teléfono y el portal antes de la presentación final.")

# Rewrite the project definition and scope to distinguish prototype from target design.
replace_paragraph(doc, "El proyecto StockFácil PC consiste", "StockFácil PC es un proyecto de gestión de inventario para una tienda de artículos de computación. El diseño contempla una aplicación móvil Android, un portal administrativo web y un servicio de datos que permite consultar y actualizar la información desde ambos canales.")
replace_paragraph(doc, "La aplicación móvil permite registrar productos utilizando", "En el prototipo actual, la aplicación móvil consulta y modifica productos, categorías y movimientos mediante el servicio de datos. La información compartida se guarda en una base SQLite de desarrollo en la computadora. El proyecto Android también contiene una tabla SQLite local para productos, pero todavía no está conectada a las pantallas; el funcionamiento sin conexión y la sincronización posterior son parte del diseño pendiente.")
replace_paragraph(doc, "La solución contempla una arquitectura de integración", "La solución usa una fuente central de información para que los registros consultados desde el teléfono y el portal correspondan al mismo inventario. En el avance local, la computadora aloja el servicio y la base de desarrollo; por ahora ambos canales requieren conectividad con ese servicio.")
replace_paragraph(doc, "El desarrollo del proyecto comprende el diseño", "El proyecto comprende la definición del proceso de inventario, los modelos de datos, los casos de uso, las reglas de negocio y los flujos de la solución, además del desarrollo e integración de las funciones previstas. Los mockups se elaborarán por separado y las evidencias se incorporarán conforme se implementen y prueben los componentes.")

# Replace scope lists that previously claimed unsupported modules.
replace_paragraph(doc, "Registro de productos desde la aplicación móvil.", "Consulta y búsqueda de productos desde la aplicación móvil.")
replace_paragraph(doc, "Almacenamiento local de productos mediante SQLite.", "Registro, edición y desactivación lógica de productos a través del servicio de datos.")
replace_paragraph(doc, "Diseño de interfaces y utilización de listas para mostrar información.", "Consulta y mantenimiento de categorías, y registro y consulta de movimientos de stock.")
replace_paragraph(doc, "Envió de productos al servidor mediante servicios REST.", "Visualización de un resumen con cantidades, alertas de stock bajo y movimientos recientes.") if any(p.text.strip().startswith("Envió de productos al servidor mediante servicios REST.") for p in doc.paragraphs) else None
replace_paragraph(doc, "Envío de productos al servidor mediante servicios REST.", "Conexión con el servicio de datos cuando el dispositivo tiene acceso a la red local.")
replace_paragraph(doc, "Gestión de productos mediante las operaciones implementadas.", "Consulta, búsqueda, registro, edición y desactivación lógica de productos.")
replace_paragraph(doc, "Administración de categorías.", "Registro, edición y desactivación de categorías; no se permite desactivar una categoría que aún tenga productos activos.")
replace_paragraph(doc, "Gestión de usuarios e inicio de sesión.", "Registro y consulta de movimientos de ingreso y salida de stock.")
replace_paragraph(doc, "Registro y consulta de movimientos de stock, incluyendo entradas y salidas.", "Resumen del inventario, productos con stock bajo y movimientos recientes.")
replace_paragraph(doc, "Consulta de reportes de inventario.", "Administración de productos y categorías desde el navegador.")
replace_paragraph(doc, "Organización de la información mediante un modelo de datos.", "Base central de desarrollo con las tablas de categorías, productos y movimientos de stock.")
replace_paragraph(doc, "Almacenamiento de productos, categorías, usuarios y movimientos de inventario, de acuerdo con la estructura implementada.", "Relaciones entre categorías, productos y movimientos, según el modelo documentado en este informe.")
replace_paragraph(doc, "Publicación de la base de datos en un servicio de hosting gratuito, conforme al requisito del proyecto.", "La publicación del servicio y la base de datos en un alojamiento externo es un objetivo pendiente.")

# Make the former limitations section follow the new design documentation.
replace_paragraph(doc, "6.3. Limitaciones del proyecto", "6.9. Limitaciones y trabajo pendiente")
replace_paragraph(doc, "El proyecto se concentra en las funcionalidades", "El prototipo funciona en una red local y requiere que la computadora mantenga iniciado el servicio para que el teléfono y el portal consulten los datos compartidos. La base de datos actual se usa en desarrollo y todavía no está publicada en un alojamiento externo.")
replace_paragraph(doc, "El funcionamiento de la aplicación móvil sin conexión", "El código de almacenamiento local en el teléfono no está conectado a las pantallas. Por ello, no se puede registrar o consultar inventario sin conexión ni sincronizar cambios pendientes. Tampoco se han incorporado inicio de sesión, gestión de usuarios, aprobación de recepción por dos responsables ni un módulo de pedidos o ventas.")
replace_paragraph(doc, "La solución se presenta como un proyecto académico", "La tienda usada como escenario no ha sido identificada ni entrevistada; el diagnóstico local, los tiempos actuales y los beneficios económicos no se han medido. Quedan por elaborar los mockups, agregar capturas de evidencia y completar la matriz de pruebas y el video de demostración.")

# Add the blueprint documents after the scope and before limitations.
anchor = find_paragraph(doc, "La publicación del servicio y la base de datos en un alojamiento externo")

anchor = insert_paragraph(doc, anchor, "6.3. Proceso de negocio propuesto", "Heading 2")
anchor = insert_paragraph(doc, anchor, "El proceso que guía StockFácil PC consiste en mantener un catálogo de artículos de computación, conocer las existencias disponibles y registrar cada ingreso o salida. El personal encargado consulta el inventario desde el teléfono y la persona responsable de la administración puede revisar los mismos datos desde el portal. Estas denominaciones representan responsabilidades del proceso; la versión actual no tiene cuentas ni permisos diferenciados por usuario.")
anchor = insert_paragraph(doc, anchor, "Casos de uso de negocio", "Heading 3")
business_cases = [
    "CB-01. Mantener el catálogo: registrar un artículo con su código, nombre, categoría, marca, modelo, precio y existencias; actualizar sus datos cuando cambien.",
    "CB-02. Consultar disponibilidad: buscar un artículo por nombre o código y revisar la cantidad disponible y el nivel mínimo definido.",
    "CB-03. Registrar una entrada: anotar el ingreso de unidades y actualizar la existencia para que el nuevo saldo pueda consultarse en ambos canales.",
    "CB-04. Registrar una salida: anotar la cantidad que sale y evitar que el saldo resultante sea negativo.",
    "CB-05. Revisar movimientos: consultar cuándo se registraron entradas y salidas, qué producto fue afectado y qué nota se dejó como referencia.",
    "CB-06. Mantener información compartida: consultar desde el portal los cambios enviados por el teléfono cuando ambos están conectados al servicio.",
    "CB-07. Continuar el trabajo sin conexión: guardar cambios en el teléfono y enviarlos al recuperar la conexión. Este caso forma parte del diseño objetivo y todavía no está implementado."
]
for text in business_cases:
    anchor = insert_paragraph(doc, anchor, text, "List Paragraph")

anchor = insert_paragraph(doc, anchor, "6.4. Documentación funcional", "Heading 2")
anchor = insert_paragraph(doc, anchor, "La siguiente especificación describe el comportamiento esperado y separa las funciones disponibles en el prototipo de las que faltan construir. Los requisitos pendientes sirven como guía para el desarrollo; no se presentan como funciones ya probadas.")
anchor = insert_paragraph(doc, anchor, "Módulos y requisitos funcionales", "Heading 3")
functional_rows = [
    ("RF-01", "Consultar el resumen del inventario: cantidad de productos, unidades, categorías, alertas de stock bajo y movimientos recientes.", "Disponible en la versión local."),
    ("RF-02", "Buscar productos por nombre, SKU, marca o modelo; el portal también permite filtrar por categoría.", "Disponible mediante el servicio."),
    ("RF-03", "Registrar, editar y desactivar productos con categoría, SKU, nombre, marca, modelo, precio, stock y stock mínimo.", "Disponible en app y portal; requiere conexión."),
    ("RF-04", "Registrar y mantener categorías; impedir su desactivación mientras existan productos activos asociados.", "Disponible en app y portal."),
    ("RF-05", "Registrar ingresos y salidas de stock y consultar el historial de movimientos.", "Disponible en app y portal; la edición/borrado de movimientos no está implementada."),
    ("RF-06", "Guardar cambios en SQLite del teléfono y sincronizarlos con el portal al recuperar la conexión.", "Pendiente de integrar en la aplicación móvil."),
    ("RF-07", "Publicar el servicio y la base de datos para que el sistema funcione fuera de la red local de desarrollo.", "Pendiente de despliegue y verificación.")
]
table = insert_table(doc, anchor, ["Código", "Función", "Estado del avance"], functional_rows, [0.65, 4.15, 1.7], 8)
anchor = insert_paragraph(doc, table, "En el teléfono, los cambios de productos, categorías y movimientos se envían al servicio cuando hay conexión. El servicio centraliza la información consultada por el portal. La tabla SQLite del dispositivo está preparada en el código, pero no interviene en estas operaciones actuales.")
anchor = insert_paragraph(doc, anchor, "6.5. Especificación de casos de uso", "Heading 2", page_break=True)
anchor = insert_paragraph(doc, anchor, "Los siguientes casos describen cómo debe interactuar el personal con el sistema. El actor indica quién realiza la tarea; no implica que el prototipo ya tenga inicio de sesión o control de permisos.")

use_cases = [
    ("CU-01. Consultar el inventario", "Actor: encargado de inventario desde Android o responsable desde el portal. Precondición: el servicio está iniciado y el dispositivo tiene conexión con la computadora.", "Flujo: 1) El usuario abre la sección de productos. 2) El sistema solicita los productos activos. 3) Presenta código, nombre, categoría, precio, cantidad y alerta cuando la cantidad es igual o menor al mínimo.", "Alternativa: si no hay conexión, se muestra un aviso y no se carga información nueva; no existe consulta offline en la versión actual."),
    ("CU-02. Buscar un producto", "Actor: encargado de inventario o responsable del portal. Precondición: existe acceso al servicio.", "Flujo: 1) El usuario escribe parte del nombre, SKU, marca o modelo. 2) El sistema filtra los productos coincidentes. 3) El usuario selecciona o revisa el resultado.", "Alternativa: si no hay coincidencias, se muestra la lista vacía; si falla la conexión, se informa el error."),
    ("CU-03. Registrar un producto", "Actor: encargado desde Android o responsable desde el portal. Precondición: hay al menos una categoría activa.", "Flujo: 1) El usuario abre Nuevo producto. 2) Completa categoría, SKU, nombre, marca, modelo, precio, stock y stock mínimo. 3) El sistema valida los datos y envía el registro. 4) El servicio guarda el producto. 5) La lista se actualiza y el registro puede consultarse desde el otro canal.", "Alternativa: se rechazan campos obligatorios vacíos, valores numéricos negativos, categoría inactiva o SKU duplicado. La aplicación requiere conexión."),
    ("CU-04. Editar un producto", "Actor: encargado desde Android o responsable del portal. Precondición: el producto está activo y existe una categoría activa.", "Flujo: 1) El usuario selecciona Editar. 2) Modifica los datos necesarios. 3) El sistema valida y guarda los cambios. 4) La lista del sistema se vuelve a consultar.", "Alternativa: un SKU que ya pertenece a otro producto, datos inválidos o un producto inexistente impiden guardar."),
    ("CU-05. Desactivar un producto", "Actor: encargado desde Android o responsable del portal. Precondición: el producto está activo.", "Flujo: 1) El usuario elige Eliminar. 2) Confirma la acción. 3) El sistema marca el producto como inactivo y deja de mostrarlo en el catálogo. 4) Los movimientos ya registrados se conservan.", "Alternativa: si el usuario cancela, no se modifica el registro. La eliminación es lógica, no borra físicamente el producto."),
    ("CU-06. Administrar categorías", "Actor: encargado desde Android o responsable del portal.", "Flujo: el usuario consulta categorías, crea una con nombre y descripción, o edita una existente. Para desactivarla, el sistema verifica que no tenga productos activos asociados.", "Alternativa: el nombre vacío o duplicado se rechaza; una categoría con productos activos no se puede desactivar."),
    ("CU-07. Registrar un movimiento de stock", "Actor: encargado desde Android o responsable del portal. Precondición: el producto está activo.", "Flujo: 1) El usuario selecciona el artículo, tipo de movimiento, cantidad y una nota opcional. 2) El sistema valida el tipo y la cantidad. 3) Para una salida, verifica que exista saldo suficiente. 4) Registra el movimiento y actualiza el stock del producto.", "Alternativa: no se acepta un tipo distinto de ingreso/salida, una cantidad menor que uno, un producto inactivo o una salida mayor al saldo disponible."),
    ("CU-08. Consultar movimientos", "Actor: encargado desde Android o responsable del portal. Precondición: el servicio está disponible.", "Flujo: 1) El usuario abre Movimientos. 2) El sistema presenta los movimientos desde el más reciente, con producto, SKU, tipo, cantidad, fecha y nota.", "Alternativa: si no hay registros, se informa que aún no existen movimientos; el historial no ofrece edición o eliminación desde la interfaz actual."),
    ("CU-09. Sincronizar cambios guardados en el teléfono", "Actor: encargado de inventario. Precondición: está implementado el almacenamiento local y vuelve la conexión.", "Flujo previsto: 1) La aplicación identifica los cambios locales pendientes. 2) Envía cada cambio al servicio. 3) El servicio valida y guarda la información. 4) La aplicación marca como sincronizados los cambios confirmados y actualiza el inventario.", "Estado: caso de uso diseñado para el alcance final; no está disponible en el prototipo. Los conflictos entre cambios simultáneos aún requieren una regla de resolución.")
]
for title, pre, flow, alternate in use_cases:
    anchor = insert_paragraph(doc, anchor, title, "Heading 3")
    anchor = insert_paragraph(doc, anchor, pre, "Normal", bold_prefix="Actor:")
    anchor = insert_paragraph(doc, anchor, flow, "Normal", bold_prefix="Flujo:")
    anchor = insert_paragraph(doc, anchor, alternate, "Normal", bold_prefix="Alternativa:")

anchor = insert_paragraph(doc, anchor, "6.6. Flujos de la aplicación", "Heading 2", page_break=True)
anchor = insert_paragraph(doc, anchor, "Flujo actual de consulta", "Heading 3")
anchor = insert_paragraph(doc, anchor, "Abrir app o portal -> solicitar productos, categorías y movimientos -> recibir datos del servicio -> mostrar resumen y listas. Si el servicio no responde, mostrar aviso de conexión y conservar la pantalla sin datos nuevos.")
anchor = insert_paragraph(doc, anchor, "Flujo actual de mantenimiento de productos", "Heading 3")
anchor = insert_paragraph(doc, anchor, "Abrir formulario -> completar datos -> validar campos, categoría, valores y SKU -> enviar al servicio -> guardar en la base central de desarrollo -> volver a consultar la lista -> mostrar el cambio en el teléfono y en el portal.")
anchor = insert_paragraph(doc, anchor, "Flujo actual de movimientos", "Heading 3")
anchor = insert_paragraph(doc, anchor, "Elegir producto -> elegir ingreso o salida -> indicar cantidad y nota opcional -> validar cantidad y saldo -> si es salida mayor al stock, rechazar -> en caso válido, guardar movimiento y actualizar existencias -> actualizar resumen e historial.")
anchor = insert_paragraph(doc, anchor, "Flujo previsto de trabajo sin conexión", "Heading 3")
anchor = insert_paragraph(doc, anchor, "Guardar el cambio en SQLite del teléfono -> dejarlo marcado como pendiente -> recuperar conexión -> enviar los cambios pendientes al servicio -> recibir confirmación -> actualizar el estado local y el inventario del portal. Este flujo es un plano de trabajo y todavía no está conectado al código de la aplicación.")

anchor = insert_paragraph(doc, anchor, "6.7. Reglas de negocio", "Heading 2", page_break=True)
rules = [
    "RN-01. Cada producto debe pertenecer a una categoría activa.",
    "RN-02. El SKU es obligatorio y no se puede repetir entre productos.",
    "RN-03. El nombre de una categoría es obligatorio y no se puede duplicar.",
    "RN-04. El precio, el stock y el stock mínimo no pueden ser negativos.",
    "RN-05. Un producto se considera con stock bajo cuando su cantidad disponible es menor o igual a su stock mínimo.",
    "RN-06. Un movimiento debe corresponder a un producto activo y su cantidad debe ser mayor que cero.",
    "RN-07. Una salida no puede superar la cantidad disponible del producto.",
    "RN-08. Todo movimiento aceptado actualiza el stock y se incorpora al historial con fecha y tipo.",
    "RN-09. Una categoría con productos activos no se puede desactivar.",
    "RN-10. Al eliminar un producto, el sistema lo desactiva y conserva sus movimientos anteriores.",
    "RN-11. Los movimientos se consultan del más reciente al más antiguo; en el prototipo no se pueden editar ni eliminar desde la interfaz.",
    "RN-12. En el prototipo, el campo de stock también puede modificarse directamente al crear o editar un producto. Ese cambio directo no crea un movimiento; por ello el historial no representa todos los ajustes de existencias.",
    "RN-13. El teléfono y el portal comparten los datos a través del servicio; el flujo actual requiere conexión. El trabajo offline y su sincronización son reglas objetivo todavía pendientes."
]
for text in rules:
    anchor = insert_paragraph(doc, anchor, text, "List Paragraph")

anchor = insert_paragraph(doc, anchor, "6.8. Modelo lógico y físico de datos", "Heading 2", page_break=True)
anchor = insert_paragraph(doc, anchor, "El modelo lógico de la base central contiene tres entidades: Categoría, Producto y Movimiento de stock. Una categoría puede organizar varios productos y un producto puede tener varios movimientos. El modelo representa la versión actual del servicio de desarrollo.")
create_er_image(ER_IMAGE)
image_paragraph = doc.add_paragraph()
image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_paragraph.add_run().add_picture(str(ER_IMAGE), width=Inches(6.35))
anchor._p.addnext(image_paragraph._p)
anchor = image_paragraph
anchor = insert_paragraph(doc, anchor, "Figura 5. Modelo entidad-relación de la base de datos central.", "Caption")
anchor = insert_paragraph(doc, anchor, "Fuente: elaboración propia a partir de las entidades y relaciones configuradas en el proyecto.", "Normal")
anchor = insert_paragraph(doc, anchor, "Modelo físico de la base central", "Heading 3")
physical_rows = [
    ("CATEGORIES", "Id INTEGER (PK); Name TEXT; Description TEXT; Active INTEGER.", "Name es único. Active indica si la categoría está disponible."),
    ("PRODUCTS", "Id INTEGER (PK); CategoryId INTEGER (FK); Sku TEXT; Name TEXT; Brand TEXT; Model TEXT; Price REAL; Stock INTEGER; MinimumStock INTEGER; Active INTEGER.", "Sku es único. CategoryId referencia CATEGORIES.Id; no se permite borrar físicamente una categoría referenciada."),
    ("STOCKMOVEMENTS", "Id INTEGER (PK); ProductId INTEGER (FK); Type TEXT; Quantity INTEGER; CreatedAt TEXT; Notes TEXT.", "ProductId referencia PRODUCTS.Id. La aplicación valida tipo y cantidad antes de guardar.")
]
table = insert_table(doc, anchor, ["Tabla", "Campos y tipos", "Relaciones y condiciones"], physical_rows, [1.2, 3.6, 1.7], 8)
anchor = insert_paragraph(doc, table, "El servicio utiliza SQLite en la computadora para la base de desarrollo. El código del proyecto Android prepara una tabla local PRODUCTS con Id, Name, Code, Price y Quantity, con código único y valores no negativos. Esa tabla local aún no está conectada a la interfaz; no debe confundirse con la base central ni considerarse evidencia de funcionamiento offline.")
anchor = insert_paragraph(doc, anchor, "Modelo local previsto para completar el objetivo", "Heading 3")
anchor = insert_paragraph(doc, anchor, "Para cumplir el guardado en el teléfono y el envío posterior, el diseño deberá conservar en el dispositivo la información necesaria para trabajar y distinguir los cambios pendientes de los ya enviados. La estructura definitiva, incluidos los cambios que ocurran en dos equipos al mismo tiempo, se cerrará al implementar y probar la sincronización; por ahora no se presenta como tabla construida.")

# Rewrite deliverables, conclusion and recommendations so they report status accurately.
replace_paragraph(doc, "El proyecto StockFácil PC contempla el desarrollo", "El avance de StockFácil PC comprende una aplicación Android, un portal administrativo web y un servicio de datos para el inventario. En esta sección se distingue lo que existe en la versión local de lo que falta integrar o documentar con evidencias.")
replace_paragraph(doc, "La aplicación móvil permite registrar productos desde un dispositivo Android.", "La aplicación móvil consulta y mantiene productos, categorías y movimientos mediante el servicio de datos cuando hay conexión. También muestra un resumen del inventario y listas para las funciones disponibles.")
replace_paragraph(doc, "Entre las funcionalidades consideradas se encuentran:", "En el prototipo móvil se encuentran disponibles las siguientes funciones:")
replace_paragraph(doc, "Registro de productos desde el dispositivo móvil.", "Consulta, búsqueda, registro, edición y desactivación lógica de productos mediante el servicio.")
replace_paragraph(doc, "Almacenamiento de los productos en la base de datos local SQLite.", "Consulta y mantenimiento de categorías y registro y consulta de movimientos de stock.")
replace_paragraph(doc, "Envío de los productos registrados hacia el servidor.", "Consulta de resumen, alertas de stock bajo y movimientos recientes.")
replace_paragraph(doc, "Visualización de las interfaces y listas necesarias para trabajar con la información del inventario.", "La tabla SQLite del teléfono está preparada en el código, pero no está conectada a estas pantallas; el modo sin conexión queda pendiente.")
replace_paragraph(doc, "Como evidencia del desarrollo, se incluirán capturas", "Las capturas de pantalla de cada operación se incorporarán en los anexos después de registrar la evidencia desde el dispositivo y el portal.")
replace_paragraph(doc, "El portal administrativo web permite gestionar", "El portal administrativo permite consultar y mantener los productos y categorías, registrar movimientos y revisar el resumen del inventario desde un navegador. En el avance actual se utiliza dentro de la red local de desarrollo.")
replace_paragraph(doc, "El portal contempla los siguientes módulos:", "Los módulos disponibles en el portal son:")
replace_paragraph(doc, "Usuarios e inicio de sesión: acceso al portal y gestión de usuarios, según las funcionalidades implementadas.", "Resumen del inventario: totales, alertas de stock bajo y movimientos recientes.")
replace_paragraph(doc, "Reportes de inventario: consulta de información relacionada con los productos y sus existencias.", "El prototipo no incluye cuentas, inicio de sesión ni administración de permisos.")
replace_paragraph(doc, "Se incorporarán capturas de pantalla de los módulos", "Las capturas de evidencia del portal y su comparación con la aplicación móvil quedan pendientes de incorporar en los anexos.")
replace_paragraph(doc, "Los servicios REST permiten la comunicación", "Los servicios de datos conectan la aplicación móvil y el portal con la base central del inventario.")
replace_paragraph(doc, "Se ha trabajado en el envío de los productos", "La aplicación consume operaciones para consultar, crear, editar y desactivar productos; consultar y mantener categorías; y registrar y consultar movimientos de stock. La prueba manual reportada en la red local confirmó que un producto registrado desde el teléfono también aparece en el portal.")
replace_paragraph(doc, "También se incluirán evidencias de las pruebas", "La matriz de pruebas con pasos, resultado esperado, resultado obtenido y captura de respaldo se completará cuando se ejecuten y documenten las pruebas funcionales de cada operación.")
replace_paragraph(doc, "La base de datos permite organizar y almacenar", "La base central de desarrollo organiza la información de categorías, productos y movimientos. El modelo entidad-relación y la descripción de campos incluidos en este informe corresponden a esa estructura.")
replace_paragraph(doc, "Este entregable comprende la documentación", "El modelo lógico, el modelo físico y el modelo entidad-relación se describen en la sección 6.8. La base central de desarrollo utiliza SQLite en la computadora; la tabla local del teléfono existe en el código, pero todavía no participa en el flujo de la aplicación.")
replace_paragraph(doc, "La publicación de la base de datos en un servicio", "La publicación del servicio y de la base de datos en un alojamiento gratuito sigue pendiente y deberá verificarse antes de afirmar que el sistema está disponible fuera de la red local.")
replace_paragraph(doc, "Se realizarán pruebas funcionales para comprobar", "La verificación manual de avance en red local confirmó que un registro creado desde el teléfono se consulta en el portal. Aún falta completar y documentar una matriz de pruebas que cubra todas las operaciones y sus respuestas ante datos inválidos o falta de conexión.")
replace_paragraph(doc, "Verificar que los productos se almacenen correctamente en SQLite.", "Comprobar la persistencia local y sincronización cuando se integre SQLite en las pantallas móviles.")
replace_paragraph(doc, "Evaluar el comportamiento de la aplicación cuando no dispone", "Evaluar el trabajo sin conexión y la recuperación de cambios pendientes cuando esa función se implemente.")
replace_paragraph(doc, "Los resultados se documentarán mediante una tabla", "Los resultados pendientes se documentarán en una tabla de casos, resultado esperado, resultado obtenido y evidencia.")
replace_paragraph(doc, "La documentación técnica reunirá los elementos", "La documentación del proyecto reúne el diseño funcional y de datos necesario para orientar el desarrollo y explicar el alcance del sistema.")
replace_paragraph(doc, "Modelo lógico y físico de la base de datos.", "Modelo lógico, modelo físico y modelo entidad-relación de la base central.")
replace_paragraph(doc, "Casos de uso y documentación funcional.", "Casos de uso de negocio, especificación de casos de uso y requisitos funcionales.")
replace_paragraph(doc, "Diagramas de flujo de la aplicación.", "Flujos actuales y flujo previsto de trabajo sin conexión y sincronización.")
replace_paragraph(doc, "Reglas de negocio.", "Reglas de negocio implementadas y reglas objetivo pendientes.")
replace_paragraph(doc, "Documentación de los servicios REST implementados.", "Descripción de los servicios de datos usados por la aplicación y el portal.")
replace_paragraph(doc, "Evidencias de las pruebas funcionales.", "Capturas y matriz de pruebas, pendientes de completar en los anexos.")
replace_paragraph(doc, "El desarrollo de StockFácil PC permite integrar", "El prototipo de StockFácil PC conecta una aplicación Android y un portal administrativo con la misma información de inventario en una red local. La consulta compartida se comprobó durante el avance; el alojamiento externo y el funcionamiento sin conexión todavía no forman parte de la versión terminada.")
replace_paragraph(doc, "La utilización de SQLite para el almacenamiento local", "La base central registra productos, categorías y movimientos, y el servicio actualiza las existencias al guardar un ingreso o una salida. La tabla SQLite del teléfono aún no está integrada; por tanto, el almacenamiento local y la sincronización no se presentan como resultados alcanzados.")
replace_paragraph(doc, "El proyecto permite aplicar los conocimientos", "La documentación funcional y el modelo de datos convierten el alcance del proyecto en una guía para completar el desarrollo. El beneficio para una tienda deberá confirmarse mediante pruebas con usuarios y mediciones del proceso antes y después de usar la solución.")
replace_paragraph(doc, "Realizar pruebas funcionales de los módulos", "Completar las pruebas de productos, categorías, movimientos y resumen; guardar capturas que muestren los resultados tanto en el teléfono como en el portal.")
replace_paragraph(doc, "Verificar la comunicación entre la aplicación móvil", "Integrar la tabla SQLite del teléfono y probar el envío de cambios pendientes cuando se recupere la conexión, incluyendo la respuesta ante un mismo producto editado desde dos equipos.")
replace_paragraph(doc, "Mantener actualizada la documentación técnica", "Publicar el servicio y la base de datos en un alojamiento aprobado para el curso, volver a probar el acceso desde el teléfono y el portal, y completar los mockups y el video demostrativo antes de la entrega final.")

# Give the inserted documentation and existing report a concise, consistent style.
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True
    if paragraph.style.name == "Normal":
        paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "Informe del proyecto StockFácil PC"
doc.core_properties.subject = "Avance y documentación funcional del proyecto de inventario"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
ER_IMAGE.unlink(missing_ok=True)
SOCIAL_CHART.unlink(missing_ok=True)
RAEE_CHART.unlink(missing_ok=True)
print(OUT)
print(f"Párrafos: {len(doc.paragraphs)} | Tablas: {len(doc.tables)} | Imágenes: {len(doc.inline_shapes)}")
