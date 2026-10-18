from docx import Document
p='outputs/StockFacil/docs/Informe_Proyecto_StockFacil_PC_Avance_reescrito.docx'
d=Document(p)
cell=d.tables[0].rows[2].cells[2]
para=cell.paragraphs[0]
para.clear()
para.add_run('Marco de referencia para administrar responsablemente la información del sistema.')
d.save(p)
print('Updated table cell')
