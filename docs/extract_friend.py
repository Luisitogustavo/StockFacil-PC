import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
p=r'C:\Users\Usuario\Downloads\Informe_S.docx'
d=Document(p)
for i, para in enumerate(d.paragraphs):
    if 81 <= i <= 258 and para.text.strip():
        print(f'[{i} {para.style.name}] {para.text}')
print('\nTABLES:')
for ti, t in enumerate(d.tables):
    print(f'-- TABLE {ti} --')
    for row in t.rows:
        print(' | '.join(cell.text.replace('\n',' / ') for cell in row.cells))
