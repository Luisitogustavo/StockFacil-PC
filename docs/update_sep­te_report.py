from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Usuario\Documents\Codex\2026-10-04\hola-me-gustar-a-saber-si\outputs\StockFacil\docs")
DOCX = ROOT / "Informe_StockFacil_PC_Avance_con_documentacion_funcional.docx"
SOCIAL_IMAGE = ROOT / "_septe_social_chart.png"
RAEE_IMAGE = ROOT / "_septe_raee_chart.png"


def find_paragraph(doc, prefix):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"No se encontró el párrafo que empieza con: {prefix}")


def replace_paragraph(doc, prefix, text):
    paragraph = find_paragraph(doc, prefix)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)
    return paragraph


def insert_after(anchor, text, style="Normal"):
    paragraph = anchor._parent.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def add_hyperlink(paragraph, label, url):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def font_file(bold=False):
    return r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"


def create_social_chart(path):
    width, height = 988, 574
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype(font_file(), 20)
    bold = ImageFont.truetype(font_file(True), 25)
    small = ImageFont.truetype(font_file(), 17)
    draw.text((width // 2, 23), "Hogares con Internet por zona, II trimestre de 2026", font=bold, fill="#17365D", anchor="mt")
    entries = [("Perú", 64.5), ("Lima Metropolitana", 84.7), ("Resto urbano", 66.2), ("Área rural", 27.1)]
    left, right, top, bottom = 260, 90, 98, 78
    plot_width = width - left - right
    plot_height = height - top - bottom
    for tick in range(0, 101, 20):
        x = left + plot_width * tick / 100
        draw.line((x, top - 5, x, height - bottom), fill="#D6DEE8", width=2)
        draw.text((x, height - bottom + 11), f"{tick}%", font=small, fill="#46586B", anchor="mt")
    row_height = plot_height / len(entries)
    for index, (label, value) in enumerate(entries):
        y = top + row_height * (index + 0.5)
        end = left + plot_width * value / 100
        draw.text((left - 14, y), label, font=regular, fill="#27384A", anchor="rm")
        draw.rounded_rectangle((left, y - 22, end, y + 22), radius=7, fill="#315F91")
        draw.text((end + 10, y), f"{value:.1f}%".replace(".", ","), font=regular, fill="#27384A", anchor="lm")
    draw.text((left + plot_width / 2, height - 21), "Porcentaje de hogares", font=small, fill="#46586B", anchor="mm")
    image.save(path, "PNG")


def create_raee_chart(path):
    width, height = 988, 574
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype(font_file(), 18)
    bold = ImageFont.truetype(font_file(True), 24)
    small = ImageFont.truetype(font_file(), 15)
    draw.text((width // 2, 22), "RAEE reciclados en el Perú, 2020–2024", font=bold, fill="#17365D", anchor="mt")
    left, right, top, bottom = 96, 34, 94, 82
    plot_width = width - left - right
    plot_height = height - top - bottom
    max_value = 30000
    for tick in range(0, max_value + 1, 5000):
        y = top + plot_height * (1 - tick / max_value)
        draw.line((left, y, width - right, y), fill="#D6DEE8", width=2)
        draw.text((left - 10, y), f"{tick:,}".replace(",", " "), font=small, fill="#46586B", anchor="rm")
    values = [5700, 11175, 14700, 20000, 25600]
    labels = ["5 700", "11 175", ">14 700", ">20 000", ">25 600"]
    years = ["2020", "2021", "2022", "2023", "2024"]
    bar_width = 96
    gap = (plot_width - len(values) * bar_width) / (len(values) + 1)
    for index, (value, label, year) in enumerate(zip(values, labels, years)):
        x0 = left + gap * (index + 1) + bar_width * index
        x1 = x0 + bar_width
        y1 = top + plot_height
        y0 = top + plot_height * (1 - value / max_value)
        draw.rounded_rectangle((x0, y0, x1, y1), radius=6, fill="#2F7D63")
        draw.text(((x0 + x1) / 2, y0 - 10), label, font=small, fill="#27384A", anchor="ms")
        draw.text(((x0 + x1) / 2, y1 + 15), year, font=regular, fill="#27384A", anchor="mt")
    draw.text((left + 4, top - 20), "Toneladas", font=small, fill="#46586B", anchor="lm")
    image.save(path, "PNG")


def replace_following_image(doc, paragraph_prefix, path):
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith(paragraph_prefix))
    for paragraph in paragraphs[start + 1:]:
        blips = paragraph._p.xpath(".//a:blip")
        if blips:
            relationship_id = blips[0].get(qn("r:embed"))
            doc.part.related_parts[relationship_id]._blob = path.read_bytes()
            return
    raise ValueError(f"No se encontró imagen asociada a: {paragraph_prefix}")


def replace_bibliography(doc):
    heading = find_paragraph(doc, "11. BIBLIOGRAFÍA")
    annex = find_paragraph(doc, "12. ANEXOS")
    body = doc._body._element
    current = False
    for element in list(body):
        if element is heading._p:
            current = True
            continue
        if element is annex._p:
            break
        if current and element.tag == qn("w:p"):
            body.remove(element)
    sources = [
        ("Instituto Nacional de Estadística e Informática (2026). El 64,5 % de los hogares del país tiene conexión a Internet durante el segundo trimestre de 2026. Resultados de la Encuesta Nacional de Hogares.", "https://www.gob.pe/institucion/inei/noticias/1450362-el-64-5-de-los-hogares-del-pais-tiene-conexion-a-internet-durante-el-segundo-trimestre-de-2026"),
        ("Ministerio de la Producción (2025a). Número de Mipyme en Perú ascendió a 2,33 millones al cierre de 2024.", "https://www.gob.pe/institucion/produce/noticias/1299194-numero-de-mipyme-en-peru-ascendio-a-2-33-millones-al-cierre-del-2024-alcanzando-un-crecimiento-de-1-6"),
        ("Ministerio de la Producción, OGEIEE (2025b). Las Mipyme en Cifras 2024. Primera edición, octubre de 2025; resultados de la Encuesta Económica Anual 2024 correspondientes al ejercicio 2023.", "https://www.producempresarial.pe/wp-content/uploads/2025/11/La-Mipyme-en-Cifras-2024_.pdf"),
        ("Presidencia del Consejo de Ministros (2023). Decreto Supremo N.° 085-2023-PCM, Política Nacional de Transformación Digital al 2030.", "https://www.gob.pe/institucion/pcm/normas-legales/4471543-085-2023-pcm"),
        ("Autoridad Nacional de Protección de Datos Personales (2025). Nuevo Reglamento de Protección de Datos Personales; vigente desde el 31 de marzo de 2025.", "https://www.gob.pe/institucion/anpd/campa%C3%B1as/128319-nuevo-reglamento-de-proteccion-de-datos-personales"),
        ("Ministerio de Justicia y Derechos Humanos (2024). Decreto Supremo N.° 016-2024-JUS, Reglamento de la Ley de Protección de Datos Personales.", "https://www.gob.pe/institucion/anpd/normas-legales/6554453-n-016-2024-jus"),
        ("Ministerio del Ambiente (2025). Minam promueve el reciclaje de residuos de aparatos eléctricos y electrónicos; cifras de Sigersol de 2020 a 2024.", "https://www.gob.pe/institucion/minam/noticias/1263803-minam-promueve-reciclaje-de-residuos-de-aparatos-electricos-y-electronicos"),
        ("Ministerio del Ambiente (2019). Decreto Supremo N.° 009-2019-MINAM, Régimen Especial de Gestión y Manejo de Residuos de Aparatos Eléctricos y Electrónicos.", "https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam"),
        ("Liu, F., Yu, Y., Fang, Y., Zhu, M., Shi, Y. y Xiao, S. (2024). Lean strategy in SMEs: Inventory leanness, operational leanness, and financial performance. The Asian Journal of Shipping and Logistics, 40(2), 109–117.", "https://doi.org/10.1016/j.ajsl.2024.02.003"),
        ("Cibertec (2026). Desarrollo de Aplicaciones Móviles I: guía del proyecto y rúbrica del curso. Material proporcionado por el estudiante.", None),
    ]
    anchor = heading
    for citation, url in sources:
        p = doc.add_paragraph(style="List Paragraph")
        p.add_run(citation + (" " if url else ""))
        if url:
            add_hyperlink(p, "Abrir fuente", url)
        p.paragraph_format.keep_together = True
        anchor._p.addnext(p._p)
        anchor = p


doc = Document(DOCX)

replace_paragraph(
    doc,
    "El acceso a Internet en los hogares peruanos varía",
    "Según los resultados de la Encuesta Nacional de Hogares publicados por el INEI (2026), durante el segundo trimestre de 2026 el 64,5 % de los hogares del Perú tenía conexión a Internet. La cifra llegó a 84,7 % en Lima Metropolitana, 66,2 % en el resto urbano y 27,1 % en el área rural. Además, el 90,2 % de las personas usuarias de Internet de seis años a más accedió mediante un teléfono celular. Estos datos describen hogares y personas a nivel nacional; no confirman la conectividad de una tienda participante.",
)
replace_paragraph(doc, "Figura 1.", "Figura 1. Hogares con acceso a Internet por zona, segundo trimestre de 2026.")
replace_paragraph(doc, "Fuente: INEI, Estadísticas de las Tecnologías", "Fuente: INEI (2026), resultados de la Encuesta Nacional de Hogares. Elaboración propia a partir del informe de tecnologías de información y comunicación en los hogares.")
replace_paragraph(
    doc,
    "La diferencia entre zonas muestra",
    "La diferencia entre zonas muestra que el acceso a Internet no es uniforme. El uso extendido del celular respalda que StockFácil PC incluya una consulta desde el teléfono, mientras que la brecha de conectividad hace relevante diseñar un proceso que no dependa de una conexión permanente. En el prototipo actual, sin embargo, tanto el teléfono como el portal requieren conexión al servicio local; el trabajo sin conexión continúa como función pendiente.",
)

replace_paragraph(
    doc,
    "El Ministerio de la Producción analizó",
    "Las Mipyme tienen un peso importante en la economía peruana: PRODUCE informó que en 2024 operaron 2 331 173 Mipyme formales, equivalentes al 99,3 % de las empresas formales, y que sus ventas representaron el 22,7 % de las ventas formales del país (PRODUCE, 2025a). En otro indicador, la Encuesta Económica Anual 2024 —con datos del ejercicio 2023— preguntó por qué algunas Mipyme no usaban plataformas digitales para vender. Entre las no usuarias, el 37,7 % dijo que no percibía mayores beneficios; el 21,4 % mencionó el precio de adquisición y el 12,4 % los costos de operación y mantenimiento. La pregunta permitía más de una respuesta (PRODUCE, 2025b).",
)
replace_paragraph(doc, "Figura 2.", "Figura 2. Motivos declarados para no usar plataformas digitales de venta, ejercicio 2023.")
replace_paragraph(doc, "Fuente: PRODUCE, Las Mipyme en Cifras 2024", "Fuente: PRODUCE–OGEIEE (2025), Las Mipyme en Cifras 2024, figura 8.13; resultados de la EEA 2024 correspondientes al ejercicio 2023. La pregunta se dirigió a Mipyme no usuarias y permitía respuestas múltiples. Elaboración propia.")
replace_paragraph(
    doc,
    "El resultado indica que el valor percibido",
    "Los datos muestran que el costo y el beneficio esperado influyen en la adopción de herramientas de venta digital; no miden el costo de un sistema de inventario ni el ahorro que produciría StockFácil PC. Un estudio de pequeñas y medianas empresas manufactureras de China también encontró que reducir inventario puede asociarse con mejores resultados hasta cierto punto, pero reducirlo en exceso puede perjudicarlos (Liu et al., 2024). Por tratarse de otro país y sector, el estudio sirve como antecedente general y no como pronóstico para una tienda peruana. El beneficio económico local deberá evaluarse midiendo el tiempo y el costo actuales de registrar, buscar y corregir existencias.",
)

replace_paragraph(
    doc,
    "El marco público peruano promueve",
    "En julio de 2023 se aprobó la Política Nacional de Transformación Digital al 2030 mediante el Decreto Supremo N.° 085-2023-PCM. La política articula la participación del Estado, el sector privado, la academia y la ciudadanía, en cuanto resulte aplicable a cada actor (PCM, 2023). Para StockFácil PC, esta norma aporta contexto sobre la dirección digital del país; no establece por sí sola que este prototipo deba incorporar funciones ajenas a la gestión de inventario.",
)
replace_paragraph(
    doc,
    "La Política Nacional de Transformación Digital",
    "El nuevo Reglamento de la Ley de Protección de Datos Personales, aprobado mediante el Decreto Supremo N.° 016-2024-JUS, está vigente desde el 31 de marzo de 2025 (ANPD, 2025). El alcance actual del prototipo comprende categorías, productos y movimientos de inventario; no registra datos personales de clientes o trabajadores ni incorpora cuentas o inicio de sesión. La norma se incluye como contexto para describir con precisión qué información maneja hoy el proyecto.",
)

replace_paragraph(
    doc,
    "La Encuesta Económica Anual 2024, cuyos resultados",
    "La Encuesta Económica Anual 2024, cuyos resultados para estos indicadores corresponden al ejercicio 2023, reportó que el 90,1 % de las Mipyme encuestadas tenía Internet, el 85,1 % contaba con teléfono móvil, el 70,9 % con computadora de escritorio y el 60,3 % con laptop. El mismo informe señaló que el 6,6 % declaró ventas por Internet y el 6,4 % compras por ese medio (PRODUCE, 2025b). Una empresa podía contar con más de un tipo de dispositivo; las cifras de equipos no se suman.",
)
replace_paragraph(doc, "Figura 3.", "Figura 3. Acceso de las Mipyme a Internet y dispositivos, ejercicio 2023.")
replace_paragraph(doc, "Fuente: PRODUCE, Las Mipyme en Cifras 2024", "Fuente: PRODUCE–OGEIEE (2025), Las Mipyme en Cifras 2024; resultados de la EEA 2024 correspondientes al ejercicio 2023. Elaboración propia.")
replace_paragraph(
    doc,
    "La presencia de teléfonos y computadoras",
    "La presencia de teléfonos y computadoras ofrece un contexto favorable para consultar información desde más de un equipo, aunque disponer de Internet no significa que la empresa ya gestione digitalmente sus existencias. La encuesta mide acceso y comercio electrónico, no el control interno del inventario. En StockFácil PC, la aplicación móvil y el portal comparten datos cuando se conectan al servicio en la red local. La tabla SQLite del teléfono todavía no está conectada a las pantallas y no se presenta como almacenamiento móvil operativo.",
)

replace_paragraph(
    doc,
    "Los artículos de computación forman parte",
    "Los artículos de computación pertenecen a una categoría de productos que, al terminar su vida útil, puede convertirse en residuos de aparatos eléctricos y electrónicos (RAEE). El MINAM señala que en el Perú se generan más de 100 000 toneladas de RAEE al año. Según los registros de Sigersol citados por el ministerio, se reciclaron 5 700 toneladas en 2020 y 11 175 en 2021; en 2022 se superaron las 14 700 toneladas, en 2023 las 20 000 y en 2024 las 25 600 (MINAM, 2025). El Decreto Supremo N.° 009-2019-MINAM establece un régimen especial para el manejo de estos residuos.",
)
replace_paragraph(doc, "Figura 4.", "Figura 4. Cantidad de RAEE reciclados en el Perú, 2020–2024.")
replace_paragraph(doc, "Fuente: MINAM, registros del Sigersol.", "Fuente: MINAM (2025), registros de Sigersol. Las etiquetas de 2022, 2023 y 2024 reflejan mínimos publicados como “más de”; no representan cifras exactas. Elaboración propia.")
replace_paragraph(
    doc,
    "Este contexto es pertinente para un negocio",
    "La información ambiental permite contextualizar la actividad comercial de artículos tecnológicos, pero no describe el manejo de residuos de una tienda participante. StockFácil PC registra productos y movimientos de existencias; no administra residuos ni acredita su reciclaje. Por ello, no se atribuye al sistema una reducción de RAEE. Si la tienda reemplaza registros en papel, ese posible efecto ambiental deberá comprobarse observando y midiendo el proceso real.",
)
replace_paragraph(
    doc,
    "La información oficial muestra",
    "La información revisada permite justificar el estudio de una herramienta que se consulte desde el teléfono y una computadora: el celular es un medio extendido de acceso, las Mipyme representan casi todas las empresas formales y la adopción de ventas por Internet aún es limitada en la encuesta citada. Las normas nacionales orientan la transformación digital y el tratamiento de datos personales, mientras que las cifras de RAEE muestran la importancia del manejo de equipos al final de su vida útil. Estos datos explican el contexto, pero no prueban una falla concreta en una tienda ni permiten prometer ahorros. Para validar la necesidad y el impacto local, se deberá observar cómo se controla hoy el inventario y medir tiempos, errores y costos con el negocio participante.",
)

for table in doc.tables:
    if table.rows and table.cell(0, 0).text.strip() == "Año" and table.cell(0, 1).text.strip() == "Marco público":
        table.cell(2, 2).text = "El prototipo actual registra datos de inventario, sin datos personales, cuentas ni inicio de sesión."
        break

create_social_chart(SOCIAL_IMAGE)
create_raee_chart(RAEE_IMAGE)
replace_following_image(doc, "Según los resultados de la Encuesta Nacional de Hogares", SOCIAL_IMAGE)
replace_following_image(doc, "Los artículos de computación pertenecen", RAEE_IMAGE)
replace_bibliography(doc)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.widow_control = True

doc.save(DOCX)
SOCIAL_IMAGE.unlink(missing_ok=True)
RAEE_IMAGE.unlink(missing_ok=True)
print(DOCX)
print(f"Párrafos: {len(doc.paragraphs)} | Tablas: {len(doc.tables)} | Figuras: {len(doc.inline_shapes)}")
