from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from pathlib import Path
p=Path('outputs/StockFacil/docs/Informe_Proyecto_StockFacil_PC_Avance_reescrito.docx')
d=Document(p)
# Repair the section heading displaced while adding the fourth summary paragraph.
para=d.paragraphs[7]
summary_text=para.text
para.clear(); para.style='Heading 1'; para.add_run('2 Introducción')
new=OxmlElement('w:p'); para._p.addprevious(new)
from docx.text.paragraph import Paragraph
np=Paragraph(new,para._parent); np.style='Normal'; np.add_run(summary_text)
# Keep source links readable and left aligned.
for para in d.paragraphs:
    if para._p.xpath('.//w:hyperlink'):
        para.alignment=WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before=Pt(2)
        para.paragraph_format.space_after=Pt(4)
        for run in para.runs: run.font.size=Pt(9)
# Keep the short political table together instead of leaving a single row on the next page.
if d.tables:
    table=d.tables[0]
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next=True
    for cell in table.rows[-1].cells:
        for para in cell.paragraphs: para.paragraph_format.keep_with_next=False

d.save(p)
print(p.resolve())
