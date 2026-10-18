from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

src = Path('outputs/StockFacil/docs/Informe_Proyecto_StockFacil_PC_Avance_SEPTE.docx')
out = Path('outputs/StockFacil/docs/Informe_Proyecto_StockFacil_PC_Avance_reescrito.docx')
doc = Document(src)

replacements = {
4: "Este informe presenta StockFácil PC, un proyecto aplicado al control de inventario de una pequeña tienda de artículos de computación. El diagnóstico organiza información social, económica, política, tecnológica y ecológica del Perú para explicar el entorno en el que se utiliza la solución y la importancia de contar con datos ordenados sobre productos y existencias.",
5: "Los objetivos se formulan con el enfoque SMART: específicos, medibles, alcanzables, relevantes y definidos en el tiempo. Se orientan a facilitar el registro y la búsqueda de productos, mantener actualizados los ingresos y salidas de stock y consultar información coincidente desde el teléfono y el portal. La mejora en los tiempos de consulta y en la calidad de los registros forma parte del propósito del proyecto y se evalúa mediante sus flujos de uso.",
6: "El avance integra una aplicación móvil, un servicio REST y un portal administrativo para trabajar con productos, categorías y movimientos de stock. La integración se comprobó en la red local al registrar un producto desde el teléfono y verificarlo en el portal. La aplicación también incorpora una estructura de almacenamiento SQLite para productos; el intercambio con el portal funciona mediante el servicio REST.",
7: "El desarrollo y la documentación emplean Android Studio para la aplicación móvil, SQLite para el almacenamiento local, servicios REST para comunicar la aplicación con el portal, Figma para los mockups e IBM Rhapsody para representar los casos de uso y los flujos. El informe reúne además los modelos lógico y físico de la base de datos, el modelo entidad-relación, las funciones y las reglas del negocio.",
9: "StockFácil PC reúne la aplicación Android, el portal administrativo y el servicio REST que comunica ambos medios. En el teléfono se consultan productos y se registran cambios de inventario; en el portal se administra el catálogo y se revisan las existencias y movimientos. La prueba realizada en la red local confirmó que un producto agregado desde el teléfono aparece también en el portal. El proyecto incorpora SQLite para el manejo local de productos y utiliza la API para compartir la información entre los dos medios.",
10: "El informe documenta el diagnóstico SEPTE, los objetivos SMART y la solución de inventario. También presenta la base de datos mediante sus modelos lógico, físico y entidad-relación; los mockups; los casos de uso; la descripción de funciones; los flujos de la aplicación; y las reglas del negocio. La propuesta se apoya en herramientas de desarrollo y modelado pertinentes para construir, explicar y revisar el sistema.",
12: "El análisis SEPTE organiza el contexto del proyecto en cinco aspectos: social, económico, político, tecnológico y ecológico. Para describirlos se utilizan estadísticas y normas de fuentes oficiales, además de una referencia internacional sobre tendencias tecnológicas. El análisis relaciona esos datos con el control de inventario de artículos de computación y con las funciones que ofrece StockFácil PC.",
14: "El uso de teléfonos móviles forma parte de la vida cotidiana de la población peruana que accede a Internet. En el segundo trimestre de 2025, el 89,0 % de las personas usuarias de Internet de seis años a más se conectó mediante un celular, con o sin plan de datos, según el INEI.",
18: "La frecuencia de acceso a Internet desde el celular respalda que el inventario se consulte desde un dispositivo móvil. StockFácil PC responde a ese hábito con una aplicación que permite revisar y actualizar productos y existencias. Las cifras describen el contexto nacional y muestran la pertinencia de ofrecer acceso móvil a la información del inventario.",
20: "La Encuesta Económica Anual 2024, con datos del ejercicio 2023, identificó motivos por los que algunas Mipyme no usaban herramientas o plataformas digitales para vender. Entre las empresas que no las utilizaban, el 37,7 % señaló que no percibía mayores beneficios; el 21,4 % mencionó el precio de adquisición y el 12,4 % los costos de operación y mantenimiento. También se registraron el uso de alternativas gratuitas (12,0 %) y el costo de capacitar al personal (11,6 %). La pregunta admitía más de una respuesta.",
24: "Los resultados muestran que el valor que una empresa reconoce en una herramienta y sus costos influyen en la adopción de canales digitales. StockFácil PC se enfoca en la gestión diaria del inventario: reúne la información de productos y existencias para facilitar su registro, búsqueda y actualización. Este enfoque permite valorar el sistema por la utilidad que aporta al trabajo con el inventario, mientras que cualquier ahorro económico se determina con mediciones de uso y costos del negocio.",
26: "En el Perú, la Política Nacional de Transformación Digital al 2030 fue aprobada mediante el Decreto Supremo N.° 085-2023-PCM. Asimismo, el Decreto Supremo N.° 016-2024-JUS aprobó el reglamento vigente de la Ley de Protección de Datos Personales. Ambas normas conforman el marco público para la transformación digital y el tratamiento responsable de la información.",
27: "StockFácil PC está orientado al registro y consulta de productos, categorías y movimientos de inventario. Por esa razón, su información principal corresponde a la operación del almacén. El análisis político presenta las normas nacionales como el marco que orienta el diseño responsable de los servicios digitales y la protección de la información que estos administran.",
30: "La Encuesta Económica Anual 2024, cuyos resultados para este indicador corresponden al ejercicio 2023, reportó que el 90,1 % de las Mipyme encuestadas contaba con Internet, el 85,1 % con teléfono móvil, el 70,9 % con computadora de escritorio y el 60,3 % con laptop. Una empresa puede disponer de más de un tipo de equipo, por lo que estos porcentajes no se suman.",
34: "La presencia de teléfonos y computadoras en las Mipyme coincide con el enfoque de StockFácil PC: una aplicación móvil para registrar y consultar el inventario y un portal administrativo para mantener el catálogo y revisar los movimientos. La aplicación incorpora SQLite para productos y se comunica con el servicio REST; la consulta del mismo producto desde el teléfono y el portal se verificó en una red local. IBM identifica tendencias tecnológicas empresariales relacionadas con la inteligencia artificial, la nube, la automatización y la ciberseguridad, que aportan un marco general para entender la evolución de los servicios digitales.",
36: "Los aparatos y componentes de computación, al terminar su vida útil, pueden convertirse en residuos de aparatos eléctricos y electrónicos (RAEE). El Ministerio del Ambiente reportó que en el Perú se valorizaron 5 724,4 toneladas de RAEE en 2020, 11 175,8 toneladas en 2021 y más de 14 700 toneladas en 2022. El Decreto Supremo N.° 009-2019-MINAM establece el régimen especial para la gestión y el manejo de estos residuos.",
40: "El volumen de RAEE valorizado en el país evidencia la importancia de gestionar adecuadamente los equipos que llegan al final de su vida útil. StockFácil PC organiza el inventario de productos y registra los movimientos de stock; la gestión de residuos electrónicos no forma parte de sus funciones. Al mantener los datos del inventario en medios digitales, el sistema también favorece que las consultas y actualizaciones del catálogo se realicen sin depender de anotaciones en papel.",
42: "Las fuentes revisadas muestran un uso extendido del celular y de herramientas digitales entre la población usuaria de Internet y las Mipyme, además de barreras de costo y percepción de utilidad que acompañan su adopción. Las normas nacionales ofrecen el marco para desarrollar servicios digitales y administrar información de manera responsable. El contexto ambiental destaca la importancia de gestionar los residuos electrónicos. En conjunto, estos aspectos dan sustento al enfoque de StockFácil PC: facilitar desde el teléfono y el portal el registro, la búsqueda y la actualización del inventario de artículos de computación.",
44: "Los objetivos describen resultados concretos del proyecto y su forma de comprobación. Se centran en el control de productos y existencias, el intercambio de información entre la aplicación y el portal, y la administración del catálogo y sus movimientos.",
46: "Mejorar el control del inventario de una pequeña tienda de artículos de computación mediante una aplicación móvil y un portal administrativo que permitan registrar, buscar y actualizar productos, conservar información de productos en SQLite y consultar los datos compartidos mediante servicios REST, con comprobaciones funcionales antes de la sustentación final del curso.",
48: "1. Implementar el registro, la búsqueda por nombre o código, la modificación y la eliminación de productos en la aplicación móvil, y conservar los datos de productos en SQLite. La comprobación consiste en ejecutar las cuatro operaciones y revisar la persistencia de los datos al cerrar y volver a abrir la aplicación antes de la sustentación final.",
49: "2. Integrar la aplicación móvil y el portal mediante servicios REST para consultar los mismos productos y movimientos de inventario. La comprobación consiste en registrar un producto desde el teléfono y verificar su aparición con el mismo código y cantidad en el portal, además de revisar el registro y consulta de movimientos.",
50: "3. Completar en el portal administrativo el mantenimiento de categorías y productos y la consulta de movimientos de stock, y documentar el modelo de datos lógico, físico y entidad-relación. La comprobación consiste en ejecutar los flujos de alta, consulta, modificación y eliminación de productos y categorías, registrar movimientos y verificar que los datos compartidos coincidan entre la aplicación y el portal antes de la sustentación final."
}

for i, text in replacements.items():
    p = doc.paragraphs[i]
    p.clear()
    p.add_run(text)

# Preserve the existing headings and chart captions, and add direct, clickable source links.
def after(paragraph, label, url):
    p = OxmlElement('w:p')
    paragraph._p.addnext(p)
    from docx.text.paragraph import Paragraph
    newp = Paragraph(p, paragraph._parent)
    newp.style = paragraph.style
    r = newp.add_run(label + ': ')
    r.font.size = None
    hyperlink = OxmlElement('w:hyperlink')
    rid = newp.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rpr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rpr.append(underline)
    run.append(rpr)
    t = OxmlElement('w:t'); t.text = url; run.append(t)
    hyperlink.append(run); newp._p.append(hyperlink)
    return newp

links = [
 (17, 'Fuente directa INEI', 'https://www.gob.pe/institucion/inei/noticias/1254582-el-60-1-de-los-hogares-tiene-conexion-a-internet'),
 (23, 'Fuente directa PRODUCE', 'https://ogeiee.produce.gob.pe/index.php/en/2017-11-05-06-56-11/estadisticas-mipyme-2/item/download/2736_eeee9391e9ed619299baa926b1d97a08'),
 (28, 'Política Nacional de Transformación Digital, PCM', 'https://www.gob.pe/institucion/pcm/normas-legales/4471543-085-2023-pcm'),
 (28, 'Reglamento de Protección de Datos Personales, MINJUSDH', 'https://www.gob.pe/institucion/anpd/normas-legales/6554453-n-016-2024-jus'),
 (33, 'Fuente directa PRODUCE', 'https://ogeiee.produce.gob.pe/index.php/en/2017-11-05-06-56-11/estadisticas-mipyme-2/item/download/2736_eeee9391e9ed619299baa926b1d97a08'),
 (33, 'Referencia de tendencias tecnológicas IBM', 'https://www.ibm.com/thought-leadership/institute-business-value/report/technology-trends'),
 (39, 'Datos de valorización de RAEE, MINAM', 'https://www.gob.pe/institucion/minam/noticias/850050-ministerio-del-ambiente-mas-de-14-700-toneladas-de-residuos-de-aparatos-electricos-y-electronicos-se-valorizan-en-el-peru'),
 (39, 'Régimen especial de RAEE, MINAM', 'https://www.gob.pe/institucion/minam/normas-legales/354138-009-2019-minam'),
]
# Insert source links after source note paragraphs, in reverse paragraph order so references stay stable.
for idx, label, url in sorted(links, key=lambda x: x[0], reverse=True):
    after(doc.paragraphs[idx], label, url)

doc.save(out)
print(out.resolve())
