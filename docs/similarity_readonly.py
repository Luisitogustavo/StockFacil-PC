import re, sys
from difflib import SequenceMatcher
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
a=Document(r'C:\Users\Usuario\Downloads\Informe_S.docx')
b=Document(r'.\outputs\StockFacil\docs\Informe_StockFacil_PC_Avance_con_documentacion_funcional.docx')
def norm(s):
    return re.sub(r'\s+',' ',re.sub(r'[^\wáéíóúñü]+',' ',s.lower(),flags=re.I)).strip()
source=[norm(p.text) for p in b.paragraphs if len(norm(p.text))>50]
friend=[p.text for p in a.paragraphs if len(norm(p.text))>50]
exact=0; near=0; bests=[]
for text in friend:
    n=norm(text)
    scores=[SequenceMatcher(None,n,m).ratio() for m in source]
    score=max(scores) if scores else 0
    if score==1: exact+=1
    if score>=0.9: near+=1
    bests.append(score)
print('Friend substantive paragraphs:',len(friend))
print('Exact matches to current working draft:',exact)
print('At least 90% text similarity:',near)
print('Median best-match similarity:',round(sorted(bests)[len(bests)//2],3))
print('Mean best-match similarity:',round(sum(bests)/len(bests),3))
print('Examples closest:')
for p in a.paragraphs:
    if len(norm(p.text))>150:
        n=norm(p.text); score=max(SequenceMatcher(None,n,m).ratio() for m in source)
        if score>0.88:
            print(round(score,3),p.text[:160])
